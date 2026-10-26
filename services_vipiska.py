"""Bank vipiskasi (Банковская выписка): parsing → dublikat tekshiruvi → tarixdan schyot → 1C qoralama.

Qo'llab-quvvatlanadigan formatlar:
  * Excel (.xlsx) — O'zbekiston banklari eksporti (sarlavha qatori avtomatik topiladi, UZ/RU ustun nomlari)
  * CSV (.csv) — ; yoki , ajratuvchi, UTF-8 / Windows-1251
  * 1CClientBankExchange (.txt) — 1C "Клиент-банк" almashinuv formati
"""
from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta

from services_1c import (ADVANCE_FOR, REPAYMENT_AUTO, E_BANK_IN, E_BANK_OUT, E_COUNTERPARTY, FIELDS, OPERATION_BY_ACCOUNT,
                         PURPOSE_KEYWORDS, ODataError, OneCService, fmt_account)

STATUS_LABELS = {
    "duplicate": "🟢 1C da bor (Dublikat - o'tkazib yuboriladi)",
    "new": "🟡 Yangi (Kiritishga tayyor)",
    "discrepancy": "🔴 Farq bor (Summa/Sana tafovuti)",
    "posted": "✅ 1C ga qoralama sifatida yozildi",
    "unchecked": "⚪ 1C da tekshirilmadi (ko'rish rejimi)",
}


class StatementError(ValueError):
    pass


@dataclass
class BankRow:
    date: date
    number: str
    direction: str  # kirim | chiqim
    amount: float
    counterparty_name: str = ""
    counterparty_inn: str = ""
    mfo: str = ""
    account_no: str = ""
    purpose: str = ""
    id: str = ""
    posted_date: date | None = None  # bank o'tkazgan sana ("Дата обработки"); bo'lmasa — date

    def key(self) -> str:
        raw = f"{self.date}|{self.number}|{self.direction}|{self.amount:.2f}|{self.counterparty_inn}"
        return "b-" + hashlib.sha1(raw.encode()).hexdigest()[:10]


@dataclass
class Statement:
    rows: list[BankRow]
    own_account: str = ""
    own_inn: str = ""  # vipiska egasi (tashkilot) STIR'i — fayl sarlavhasidan yoki hujjatlardan
    own_name: str = ""  # vipiska egasi nomi (masalan "Наименование счёта ООО ...")
    source: str = ""
    warnings: list[str] = field(default_factory=list)
    is_sample: bool = False  # namuna (test) vipiska — jonli 1C ga hech qachon yozilmaydi
    opening_balance: float | None = None  # davr boshidagi qoldiq (vipiska sarlavhasidan), bo'lmasa — None

    @property
    def closing_balance(self) -> float | None:
        if self.opening_balance is None:
            return None
        return round(self.opening_balance + sum(r.amount if r.direction == "kirim" else -r.amount for r in self.rows), 2)

    @property
    def currency_code(self) -> str:
        """O'zbekiston hisob raqamida 6–8-xonalar valyuta kodi: 000 = so'm, 840 = USD, 978 = EUR …"""
        return self.own_account[5:8] if len(self.own_account) == 20 else ""


# ---------------------------------------------------------------- qiymatlarni normallashtirish
def parse_amount(v) -> float | None:
    """Summani o'qish. None — o'qib bo'lmadi (qator ogohlantirish bilan o'tkazib yuboriladi).

    O'nlik ajratuvchi — oxirgi "." yoki "," va undan keyin 1–2 xona (tiyin). Undan keyin 3 xona bo'lsa,
    bu minglik ajratuvchi: "1.250.000" = 1 250 000, "1,250" = 1 250, "1.250.000,50" = 1 250 000.50.
    """
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip()
    neg = s.startswith("-") or s.endswith("-") or (s.startswith("(") and s.endswith(")"))
    s = re.sub(r"[^0-9.,]", "", s)
    if not re.search(r"\d", s):
        return None
    m = re.search(r"[.,](\d{1,2})$", s)
    if m:
        whole, frac = s[:m.start()], m.group(1)
    else:
        whole, frac = s, ""
    whole = re.sub(r"[.,]", "", whole)
    try:
        value = float(f"{whole or '0'}.{frac or '0'}")
    except ValueError:
        return None
    return round(-value if neg else value, 2)


_DIR_IN_EXACT = {"k", "к", "c", "cr", "кт", "+"}
_DIR_OUT_EXACT = {"d", "д", "dr", "дт", "-"}
_DIR_IN_WORDS = ("kirim", "kredit", "кредит", "приход", "поступ", "credit", "tushum")
_DIR_OUT_WORDS = ("chiqim", "debet", "дебет", "расход", "спис", "debit", "to'lov")


def parse_direction(v) -> str | None:
    t = str(v or "").strip().lower()
    if t in _DIR_IN_EXACT or any(w in t for w in _DIR_IN_WORDS):
        return "kirim"
    if t in _DIR_OUT_EXACT or any(w in t for w in _DIR_OUT_WORDS):
        return "chiqim"
    return None


def parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # Excel seriya raqami
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip().split(" ")[0]
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%Y.%m.%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def clean_code(v) -> str:
    """STIR / MFO: faqat raqamlar (Excel float formatini ham to'g'rilaydi)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\D", "", str(v))


def clean_account(v) -> tuple[str, bool]:
    """Hisob raqami (20 xona). Excel'da SON sifatida saqlangan bo'lsa, 15 xonadan keyingi raqamlar allaqachon
    yo'qolgan — bunday raqamni ishlatmaymiz. Natija: (raqam, aniqlik_yo'qolganmi)."""
    if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) >= 1e15:
        return "", True
    return clean_code(v), False


SEQ_HEADER = re.compile(r"^(№|n|#)?\s*(пп|п/п|п\.п\.?|t/r|т/р|tr)$")


def classify_header(h) -> str | None:
    s = " ".join(str(h or "").lower().split())  # "Обороты по\nдебету" -> "обороты по дебету"
    if not s:
        return None
    if SEQ_HEADER.match(s):
        return "seq"  # tartib raqami ("№ пп") — hujjat raqami emas
    rules = [
        ("purpose", ("maqsad", "назначен", "purpose", "izoh")),
        ("inn", ("stir", "inn", "инн")),
        ("mfo", ("mfo", "мфо", "бик")),
        ("debit", ("debet", "дебет", "chiqim", "расход", "списан")),
        ("credit", ("kredit", "кредит", "kirim", "приход", "поступ")),
        ("name", ("наимен", "nomi")),  # "Наименование счёта" — bu kontragent NOMI, hisob raqami emas
        ("account", ("hisob", "счет", "счёт", "account")),
        ("name", ("kontragent", "контрагент", "korrespondent", "корреспондент", "получател", "плательщ")),
        ("date_posted", ("дата обработ", "дата провод", "дата исполн", "o'tkazilgan sana", "дата операц")),
        ("date", ("sana", "дата", "date")),
        ("number", ("hujjat", "№", "номер", "nomer", "док", "doc")),
        ("amount", ("summa", "сумма")),
        ("direction", ("turi", "тип", "д/к", "вид")),
    ]
    for name, keys in rules:
        if any(k in s for k in keys):
            return name
    return None


# ---------------------------------------------------------------- parserlar
OPENING_RE = re.compile(r"начало дня|входящ\w*\s+(?:сальдо|остаток)|начальн\w*\s+остаток|boshlang.ich qoldiq|"
                        r"kirish qoldig", re.IGNORECASE)


def _balance_in(row: list, text: str) -> float | None:
    """Qoldiq qatoridagi son: alohida katakdagi raqam yoki matndagi "…: 1 000 000,00"."""
    for c in row:
        if isinstance(c, (int, float)) and not isinstance(c, bool):
            return float(c)
    m = re.search(r"[:=]\s*(-?[\d\s.,]+\d)", text)
    return parse_amount(m.group(1)) if m else None


def _rows_from_table(table: list[list]) -> Statement:
    header_idx, mapping = None, {}
    for i, row in enumerate(table[:40]):
        m = {}
        for j, cell in enumerate(row):
            kind = classify_header(cell)
            if kind and kind not in m:
                m[kind] = j
        if "date" not in m and "date_posted" in m:
            m["date"] = m["date_posted"]  # yagona sana ustuni — o'tkazilgan sana
        if "date" in m and ({"debit", "credit"} <= m.keys() or "amount" in m) and len(m) >= 4:
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise StatementError("Vipiska sarlavha qatori topilmadi. Kerakli ustunlar: Sana, Hujjat №, "
                             "Debet/Kredit (yoki Summa), Kontragent, STIR, To'lov maqsadi.")
    st = Statement(rows=[])
    for k in range(header_idx):  # sarlavha ustidagi qatorlardan o'z hisob raqami va STIR'ini topish
        text = " ".join(str(c) for c in table[k] if c is not None)
        m = re.search(r"(\d{20})", text)
        if m and not st.own_account:
            st.own_account = m.group(1)
        m = re.search(r"(?:STIR|ИНН|INN)\D{0,6}(\d{9})\b", text, re.IGNORECASE)
        if m and not st.own_inn:
            st.own_inn = m.group(1)
        m = re.search(r"(?:наименование сч[её]та|наименование клиента|mijoz|hisob nomi)[:\s]+(.+)", text, re.IGNORECASE)
        if m and not st.own_name:
            st.own_name = m.group(1).strip()
        if st.opening_balance is None and OPENING_RE.search(text):
            st.opening_balance = _balance_in(table[k], text)

    def cell(row, kind):
        j = mapping.get(kind)
        return row[j] if j is not None and j < len(row) else None

    def text_of(row) -> str:
        return " ".join(str(c) for c in row if c not in (None, ""))[:60]

    for offset, row in enumerate(table[header_idx + 1:], start=header_idx + 2):
        if not any(c not in (None, "") for c in row):
            continue
        if any(str(c or "").strip().lower().startswith(("jami", "итого", "всего", "оборот"))
               for j, c in enumerate(row) if j != mapping.get("purpose")):
            break
        if all(cell(row, k) in (None, "", 0) for k in ("debit", "credit", "amount")):
            continue  # summa yo'q — tranzaksiya emas (sarlavha, izoh, qoldiq qatori)
        raw_date = cell(row, "date")
        d = parse_date(raw_date)
        if not d:
            low = text_of(row).lower()
            if not any(w in low for w in ("остаток", "qoldiq", "saldo", "сальдо")):
                st.warnings.append(f"{offset}-qator: sana o'qilmadi ({raw_date!r}) — qator o'tkazib yuborildi: {text_of(row)}")
            continue
        if "debit" in mapping or "credit" in mapping:
            debit, credit = parse_amount(cell(row, "debit")), parse_amount(cell(row, "credit"))
            if debit is None or credit is None:
                st.warnings.append(f"{offset}-qator: summa o'qilmadi — qator o'tkazib yuborildi: {text_of(row)}")
                continue
            if debit and credit:
                st.warnings.append(f"{offset}-qator: bir qatorda ham debet, ham kredit — qator o'tkazib yuborildi")
                continue
            direction, amount = ("chiqim", abs(debit)) if debit else ("kirim", abs(credit))
        else:
            amount = parse_amount(cell(row, "amount"))
            if amount is None:
                st.warnings.append(f"{offset}-qator: summa o'qilmadi — qator o'tkazib yuborildi: {text_of(row)}")
                continue
            if "direction" in mapping:
                direction = parse_direction(cell(row, "direction"))
                if not direction:
                    st.warnings.append(f"{offset}-qator: turi aniqlanmadi ({cell(row, 'direction')!r}) — "
                                       "qator o'tkazib yuborildi")
                    continue
            else:
                direction = "chiqim" if amount < 0 else "kirim"
            amount = abs(amount)
        if not amount:
            st.warnings.append(f"{offset}-qator: summa 0 — qator o'tkazib yuborildi: {text_of(row)}")
            continue
        num = cell(row, "number")
        if isinstance(num, float) and num.is_integer():
            num = int(num)
        account_no, lost = clean_account(cell(row, "account"))
        if lost:
            st.warnings.append(f"{offset}-qator: hisob raqami Excel'da SON formatida — oxirgi raqamlari yo'qolgan, "
                               "e'tiborga olinmadi. Vipiskani hisob raqami 'Matn' formatida bo'lgan holda eksport qiling.")
        posted = parse_date(cell(row, "date_posted")) if "date_posted" in mapping else None
        st.rows.append(BankRow(
            date=d, posted_date=posted if posted and posted != d else None,
            number=str(num or "").strip(), direction=direction, amount=amount,
            counterparty_name=str(cell(row, "name") or "").strip(), counterparty_inn=clean_code(cell(row, "inn")),
            mfo=clean_code(cell(row, "mfo")), account_no=account_no,
            purpose=str(cell(row, "purpose") or "").strip()))
    return st


def parse_excel(content: bytes) -> Statement:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise StatementError(f"Excel faylni o'qib bo'lmadi ({e}). Eski .xls bo'lsa — .xlsx ga saqlang.") from e
    last_err = None
    for ws in wb.worksheets:
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        try:
            return _rows_from_table(table)
        except StatementError as e:
            last_err = e
    raise last_err or StatementError("Excel faylda varaq topilmadi")


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def parse_csv(content: bytes) -> Statement:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        class dialect(csv.excel):  # noqa: N801 — global csv.excel ni o'zgartirmaymiz
            delimiter = ";"
    return _rows_from_table([list(r) for r in csv.reader(io.StringIO(text), dialect)])


def parse_1c_exchange(content: bytes) -> Statement:
    text = _decode(content)
    if "1CClientBankExchange" not in text[:200]:
        raise StatementError("Fayl 1CClientBankExchange formatida emas")
    st = Statement(rows=[])
    doc: dict | None = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line or "=" not in line and line != "КонецДокумента":
            continue
        if line.startswith("РасчСчет=") and not st.own_account:
            st.own_account = clean_code(line.split("=", 1)[1])
        elif line.startswith("НачальныйОстаток=") and st.opening_balance is None:
            st.opening_balance = parse_amount(line.split("=", 1)[1])
        elif line.startswith("СекцияДокумент="):
            doc = {}
        elif line == "КонецДокумента" and doc is not None:
            if st.own_account and not st.own_inn:
                for side in ("Плательщик", "Получатель"):
                    if clean_code(doc.get(f"{side}Счет") or doc.get(f"{side}РасчСчет")) == st.own_account:
                        st.own_inn = clean_code(doc.get(f"{side}ИНН"))
            row, n = _exchange_doc(doc, st.own_account), doc.get("Номер", "?")
            if row is None:
                st.warnings.append(f"№{n}: sana o'qilmadi ({doc.get('Дата')!r}) — hujjat o'tkazib yuborildi")
            elif not row.amount:
                st.warnings.append(f"№{n}: summa o'qilmadi yoki 0 ({doc.get('Сумма')!r}) — hujjat o'tkazib yuborildi")
            else:
                if not st.own_account and not (doc.get("ДатаСписано") or doc.get("ДатаПоступило")):
                    st.warnings.append(f"№{n}: yo'nalish (kirim/chiqim) aniq emas — 'kirim' deb olindi, tekshiring")
                st.rows.append(row)
            doc = None
        elif doc is not None:
            k, v = line.split("=", 1)
            doc[k.strip()] = v.strip()
    if not st.rows:
        raise StatementError("1CClientBankExchange faylida hujjat topilmadi")
    return st


def _exchange_doc(d: dict, own: str) -> BankRow | None:
    dt = parse_date(d.get("Дата"))
    if not dt:
        return None
    payer_acc = clean_code(d.get("ПлательщикСчет") or d.get("ПлательщикРасчСчет"))
    if own and payer_acc == own:
        out = True
    elif own and clean_code(d.get("ПолучательСчет") or d.get("ПолучательРасчСчет")) == own:
        out = False
    else:
        out = bool(d.get("ДатаСписано")) and not d.get("ДатаПоступило")
    side = "Получатель" if out else "Плательщик"
    return BankRow(
        date=dt, number=d.get("Номер", ""), direction="chiqim" if out else "kirim", amount=abs(parse_amount(d.get("Сумма")) or 0),
        counterparty_name=d.get(f"{side}1") or d.get(side, ""), counterparty_inn=clean_code(d.get(f"{side}ИНН")),
        mfo=clean_code(d.get(f"{side}МФО") or d.get(f"{side}БИК")),
        account_no=clean_code(d.get(f"{side}Счет") or d.get(f"{side}РасчСчет")),
        purpose=d.get("НазначениеПлатежа", ""))


def parse_statement(filename: str, content: bytes) -> Statement:
    name = (filename or "").lower()
    if len(content) > 10 * 1024 * 1024:
        raise StatementError("Fayl juda katta (10 MB dan oshmasin)")
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        st = parse_excel(content)
    elif name.endswith(".txt") or b"1CClientBankExchange" in content[:200]:
        st = parse_1c_exchange(content)
    elif name.endswith(".csv"):
        st = parse_csv(content)
    elif name.endswith(".xls"):
        raise StatementError("Eski .xls format qo'llab-quvvatlanmaydi — faylni Excel'da .xlsx sifatida saqlang")
    else:
        raise StatementError("Format noma'lum. .xlsx, .csv yoki 1CClientBankExchange .txt yuklang")
    if not st.rows:
        detail = (" Sabablar: " + "; ".join(st.warnings[:5])) if st.warnings else ""
        raise StatementError("Vipiskada birorta ham tranzaksiya topilmadi." + detail)
    if st.currency_code and st.currency_code != "000":
        st.warnings.insert(0, f"Hisob raqami valyuta hisobiga o'xshaydi (kod {st.currency_code}). Hozircha faqat so'm "
                              "(UZS) hisoblari 1C ga yoziladi — qatorlar faqat ko'rish uchun.")
    seen: dict[str, int] = {}
    for r in st.rows:
        base = r.key()
        seen[base] = seen.get(base, 0) + 1
        r.id = base if seen[base] == 1 else f"{base}-{seen[base]}"
    st.source = filename
    return st


# ---------------------------------------------------------------- tahlil
# ---------------------------------------------------------------- shartnoma raqami (to'lov maqsadidan)
_KW = r"(?:договор\w*|дог\.?|shartnoma\w*|шартнома\w*|контракт\w*)"
_NUM = r"(?:№|N_?|N\s|#)?\s*([A-ZА-ЯЁ]?-?\d[\w/-]*)"
_DATE = r"(\d{2}\.\d{2}\.\d{4})"
_SONLI = r"(?:\s*-?\s*(?:сонли|sonli|сон|son))?"
CONTRACT_PATTERNS = [
    # 11.10.2026 yildagi 398 shartnomaga / 11.10.2026ЙИЛДАГИ №399 СОНЛИ ШАРТНОМАГА / 11.09.2026 yildagi 240 sonli shartnoma
    re.compile(_DATE + r"\s*(?:й|y)\w*\s*" + _NUM + _SONLI + r"\s*" + _KW, re.IGNORECASE),
    # дог№50 от 10.09.2026 / Договор 439 от 16.10.2026 / Shartnoma № A-12/3 dan 01.02.2026 / Shartnoma N_434
    re.compile(_KW + r"\s*" + _NUM + r"(?:\s*(?:от|dan|дан)\s*" + _DATE + r")?", re.IGNORECASE),
    # №440 шартномага / 17-sonli shartnoma / 05/26-sonli shartnoma
    re.compile(r"(?:№\s*)?(\d[\w/]*)" + _SONLI + r"\s*" + _KW, re.IGNORECASE),
]


def extract_contract(purpose: str) -> dict | None:
    """To'lov maqsadidagi shartnoma: {"number": "398", "date": "2026-10-11" | None, "text": topilgan qism}."""
    text = " ".join(str(purpose or "").split())
    for i, pat in enumerate(CONTRACT_PATTERNS):
        m = pat.search(text)
        if not m:
            continue
        if i == 0:
            raw_date, number = m.group(1), m.group(2)
        elif i == 1:
            number, raw_date = m.group(1), m.group(2)
            if not raw_date:  # "15.10.2026yil Shartnoma N_434" — sana kalit so'zdan oldin
                before = re.search(_DATE + r"\s*(?:й|y)\w*\s*$", text[:m.start()], re.IGNORECASE)
                raw_date = before.group(1) if before else None
        else:
            number, raw_date = m.group(1), None
        number = number.strip(".,;:-")
        if not number or not re.search(r"\d", number) or len(number) > 20:
            continue
        d = parse_date(raw_date) if raw_date else None
        return {"number": number, "date": d.isoformat() if d else None, "text": m.group(0)}
    return None


TAX_INNS = {"201122919"}  # Moliya vazirligi huzuridagi G'aznachilik (soliq to'lovlari); demo STIR bu yerga qo'shilmaydi
TAX_NAME = re.compile(r"молия вазирлиги|moliya vazirligi|ғазна|газна|g'azna|gazna|казначейств|\bдси\b|\bdsi\b|"
                      r"солиқ|soliq|налогов", re.IGNORECASE)
BANK_FEE = re.compile(r"за документ|комисс|komissiya|за обслуживан|xizmat haqi|за услуги банка|bank xizmat",
                      re.IGNORECASE)


# ---------------------------------------------------------------- aniq belgilar (soliq kodi, ish haqi, kredit …)
# Bular 1C tarixidan USTUN: G'aznachilik bitta kontragent, lekin har bir soliq o'z schyotiga olinadi.
# Soliq kodlari: G'aznachilik to'lov maqsadi formati "~hisob~STIR~kod ..." yoki maqsad boshida kod ("101 …").
TAX_CODE_ACCOUNTS = {
    "36": ("65.20", "ijtimoiy soliq (kod 36)"),
    "46": ("64.10", "JShDS — jismoniy shaxslar daromad solig'i (kod 46)"),
    "101": ("65.20", "jamg'arib boriladigan pensiya badali / INPS (kod 101)"),
}
QQS_RE = re.compile(r"\bqqs\b|\bндс\b|\bnds\b|ққс|ққс", re.IGNORECASE)  # faqat soliq organiga to'lovda
TAX_CODE_RE = re.compile(r"~\d{9,}~\d{9}~(\d{1,3})\b|^\s*(?:\d{5}\s+)?(\d{2,3})\s+\D")  # "08102 101 …" ham
CLASSIFIERS = [  # (regex, schyot, izoh, yo'nalishlar)
    (r"ижтимоий|ijtimoiy soliq|соц\.?\s*налог|социальн\w*\s+налог|yagona ijtimoiy|\bесп\b", "65.20", "ijtimoiy soliq",
     ("chiqim",)),
    (r"жисмоний шахслар даромад|jismoniy shaxslar daromad|\bjshds\b|\bжшдс\b|\bндфл\b|daromad solig", "64.10",
     "JShDS (daromad solig'i)", ("chiqim",)),
    (r"пенсия бадал|pensiya badal|\bинпс\b|\binps\b|жам.ариб бориладиган|jamg.arib boriladigan", "65.20",
     "pensiya badali / INPS", ("chiqim",)),
    (r"ish haqi|иш ҳақи|иш хаки|\boylik\b|зарплат|заработн\w* плат|аванс по з/п", "67.10", "ish haqi", ("chiqim",)),
    (r"dividend|дивиденд", "66.10", "dividend", ("chiqim",)),
    (r"kredit foiz|процент\w* по кредит|кредит\w* бўйича фоиз|foiz to.lov", "96.10", "kredit foizi", ("chiqim",)),
    (r"погашени\w* кредит|kredit asosiy|основно\w* долг|молиявий ёрдам|moliyaviy yordam|финансов\w* помощ|фин\.?\s*помощ"
     r"|\bзай[мё]", "68.10", "kredit / moliyaviy yordam", ("chiqim", "kirim")),
    (r"korporativ karta|корпоративн\w* карт|\bкпк\b|\bkpk\b", "55.20", "korporativ karta (KPK)", ("chiqim",)),
    (r"terminal|uzcard|\bhumo\b|эквайринг|ekvayring|inkassats|инкассац|\bclick\b|\bpayme\b", "57.10",
     "terminal / inkassatsiya tushumi", ("kirim",)),
]
CLASSIFIERS = [(re.compile(p, re.IGNORECASE), acc, why, dirs) for p, acc, why, dirs in CLASSIFIERS]


def is_tax_payee(r: BankRow) -> bool:
    return r.counterparty_inn in TAX_INNS or bool(TAX_NAME.search(r.counterparty_name or ""))


def tax_code_of(purpose: str) -> str | None:
    """G'aznachilik to'lovidagi soliq kodi: "~hisob~STIR~46 …" yoki "101 …" -> "46" / "101"."""
    m = TAX_CODE_RE.search(" ".join(str(purpose or "").split()))
    return (m.group(1) or m.group(2)) if m else None


async def tax_code_history(onec: OneCService, entity: str, kontr_key: str | None, code: str | None) -> dict | None:
    """Shu soliq kodi bilan firmadagi to'lovlar (BARCHA soliq to'lovlari, qaysi soliq organiga yozilganidan qat'i nazar):
    eng oxirgisi — namuna (soliq turi Налог_Key, kontragent, schyot provodkadan). 1C da soliq to'lovining schyotini
    SOLIQ TURI belgilaydi; tarixda bir kod uchun turli soliq turi/schyot bo'lsa — ogohlantirish (adashib tanlangan
    bo'lishi mumkin)."""
    if not code or onec.mode in ("offline", "file"):
        return None
    from services_1c import E_TAXES, _has_lines, _is_empty_key
    docs = [d for d in await onec.tax_docs(entity) if tax_code_of(d.get(FIELDS["purpose"])) == code]
    if not docs:
        return None
    posted = await onec.bank_postings(entity, datetime.combine(date.today() - timedelta(days=400), datetime.min.time()))
    history = []  # (sana, schyot, soliq turi, hujjat) — yangisi birinchi
    for d in sorted(docs, key=lambda x: x.get("Date") or "", reverse=True):
        key = posted.get(d.get("Ref_Key")) if not _has_lines(d) else None
        key = key or OneCService._extract_pattern(d)["account_key"]
        acc = await onec.account_code(key) if key else None
        if acc:
            tax = d.get("Налог_Key")
            history.append((str(d.get("Date") or "")[:10], acc, None if _is_empty_key(tax) else tax, d))
    if not history:
        return None
    names = await onec.describe(E_TAXES, [t for _d, _a, t, _doc in history if t])
    tax_name = lambda t: (names.get(t) or {}).get("Description") or "—"  # noqa: E731
    last_day, account, tax_key, doc = history[0]
    when = f"{last_day[8:10]}.{last_day[5:7]}.{last_day[:4]}"
    others = Counter((tax_name(t), a) for _d, a, t, _doc in history[1:] if (t, a) != (tax_key, account))
    note = (" (oldin: " + ", ".join(f"{a}×{n}" for (_t, a), n in others.most_common(2)) + ")") if others else ""
    warning = None
    if others:
        (o_name, o_acc), o_n = others.most_common(1)[0]
        warning = (f"⚠️ Soliq kodi {code}: oxirgi to'lov ({when}) 1C da «{tax_name(tax_key)}» turi bilan → {account}, "
                   f"oldingi {o_n} tasida «{o_name}» → {o_acc}. Qaysi biri to'g'ri ekanini tekshiring.")
    return {"account": account, "hits": sum(1 for _d, a, _t, _x in history if a == account), "total": len(history),
            "tax_key": tax_key, "tax_name": tax_name(tax_key) if tax_key else "", "template_ref": doc.get("Ref_Key"),
            "counterparty_key": OneCService.counterparty_of(doc), "warning": warning,
            "reason": f"🧾 Soliq kodi {code}: shu firmadagi oxirgi to'lov {when} → {account}"
                      + (f" («{tax_name(tax_key)}»)" if tax_key else "") + note}


def classify_purpose(r: BankRow) -> dict | None:
    """Aniq belgi bo'yicha schyot: {"account", "reason"} yoki None (unda 1C tarixi / qoidalar / standart)."""
    purpose = " ".join(str(r.purpose or "").split())
    if r.direction == "chiqim" and is_tax_payee(r):
        m = TAX_CODE_RE.search(purpose)
        code = m and (m.group(1) or m.group(2))
        if code in TAX_CODE_ACCOUNTS:
            acc, why = TAX_CODE_ACCOUNTS[code]
            return {"account": acc, "reason": f"🧾 Soliq kodi: {why} → {acc}"}
        if QQS_RE.search(purpose):
            return {"account": "64.10", "reason": "🧾 Soliq organiga QQS to'lovi → 64.10"}
    for rx, acc, why, dirs in CLASSIFIERS:
        if r.direction in dirs and rx.search(purpose):
            return {"account": acc, "reason": f"🧾 To'lov maqsadi: {why} → {acc}"}
    return None  # oddiy G'aznachilik to'lovi (kodsiz) — 1C tarixi aniqroq; bo'lmasa standart 64.10


def rule_account(r: BankRow) -> str:
    """1C tarixi va qoidalar bo'lmasa — standart taxmin."""
    classified = classify_purpose(r)
    if classified:
        return classified["account"]
    low = (r.purpose or "").lower()
    if r.direction == "chiqim":
        if r.counterparty_inn in TAX_INNS or TAX_NAME.search(r.counterparty_name or ""):
            return "64.10"
        if (r.account_no or "").startswith("164") or BANK_FEE.search(low):
            return "94.20"
    for kw, acc, kw_dir in PURPOSE_KEYWORDS:
        if kw in low and kw_dir == r.direction:
            return acc
    return "60.10" if r.direction == "chiqim" else "40.10"


def operation_for(account: str, direction: str) -> str:
    ops = OPERATION_BY_ACCOUNT.get(account)
    if not ops:
        return "ПрочееСписание" if direction == "chiqim" else "ПрочееПоступление"
    return ops[0] if direction == "chiqim" else ops[1]


SALARY_WORDS = ("ish haqi", "ish xaqi", "иш ҳақи", "зарплат", "заработн", "oylik maosh", "plastik kartalarga",
                "пластиковые карты")


def special_kind(r: BankRow, org_inn: str) -> tuple[str | None, str | None]:
    """Avtomatik yozilmaydigan maxsus operatsiyalar: (turi, to'siq matni)."""
    if org_inn and r.counterparty_inn == org_inn and (
            (r.account_no or "").startswith("164") or BANK_FEE.search(r.purpose or "")):
        return "bank_fee", None  # bank o'z xizmat haqini yechgan (qatorda firmaning o'z STIR'i turadi)
    if org_inn and r.counterparty_inn == org_inn:
        return "transfer", ("🔁 O'z hisoblaringiz orasidagi o'tkazma — 1C da \"Перевод на другой счет\" sifatida "
                            "qo'lda kiriting (avtomatik kiritish qo'llab-quvvatlanmaydi)")
    low = (r.purpose or "").lower()
    if r.direction == "chiqim" and any(w in low for w in SALARY_WORDS):
        return "salary", ("💳 Ish haqi to'lovi — 1C da \"Перечисление заработной платы\" (vedomost bilan) orqali "
                          "qo'lda kiriting")
    return None, None


async def resolve_contract(row: dict, onec: OneCService) -> dict:
    """Qaysi shartnoma: to'lov maqsadida raqam bo'lsa — aynan o'sha (1C dan), aks holda keyinroq tarixdan/standart."""
    ref = row.get("contract_ref")
    if not ref:
        return {"key": None, "source": None, "label": ""}
    when = f" ({ref['date'][8:10]}.{ref['date'][5:7]}.{ref['date'][:4]})" if ref.get("date") else ""
    if not row.get("counterparty_key"):
        return {"key": None, "source": "unchecked", "label": f"📄 Shartnoma №{ref['number']}{when} — 1C da tekshirilmadi"}
    contract = await onec.find_contract_by_ref(row["counterparty_key"], ref["number"], ref.get("date"))
    if contract:
        return {"key": contract["Ref_Key"], "source": "purpose", "name": contract.get("Description", ""),
                "label": f"📄 Shartnoma №{ref['number']}{when} — 1C da topildi ✓"}
    # 1C da yo'q — yangi shartnoma bo'yicha birinchi to'lov: 1C ga kiritishda shartnoma yaratiladi
    return {"key": None, "source": "new", "number": ref["number"], "date": ref.get("date"),
            "label": f"➕ Yangi shartnoma №{ref['number']}{when} — 1C ga kiritishda yaratiladi"}


def rule_brief(rule: dict) -> dict:
    return {k: rule.get(k) for k in ("id", "kind", "match", "account", "advance_account", "operation", "contract_key",
                                     "cash_flow_key", "locked")}


def advance_for(row: dict) -> str:
    """Qarz bo'lmasa ishlatiladigan avans schyoti: qoida → 1C tarixi → standart juftlik (40.10→63.10, 60.10→43.10)."""
    acc = row.get("account") or ""
    rule, learned = row.get("rule") or {}, row.get("learned") or {}
    if row.get("account_source") in ("rule", "keyword") and rule.get("account") == acc and rule.get("advance_account"):
        return rule["advance_account"]
    if learned.get("account_code") == acc and learned.get("advance_code"):
        return learned["advance_code"]
    if rule.get("account") == acc and rule.get("advance_account"):
        return rule["advance_account"]
    return ADVANCE_FOR.get(acc, "")


def _sum(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ")


def settlement_note(row: dict) -> str:
    acc, adv = row.get("account"), row.get("advance_account")
    if not adv:
        return ""
    split = row.get("balance_split") or {}
    if split.get("account") == acc and split.get("found"):
        when = f"{split['as_of'][8:10]}.{split['as_of'][5:7]}.{split['as_of'][:4]} " if split.get("as_of") else ""
        if split["debt"] <= 0:
            return f"💡 OSV {when}bo'yicha qarz yo'q — butun summa {adv} (avans) ga tushadi"
        if split["advance"] <= 0:
            return f"💡 OSV {when}bo'yicha qarz {_sum(split['debt'])} — to'lov to'liq {acc} qarzini yopadi"
        return (f"💡 OSV {when}bo'yicha qarz {_sum(split['debt'])}: {_sum(split['close'])} → {acc} (qarz yopiladi), "
                f"{_sum(split['advance'])} → {adv} (avans)")
    if row.get("kind") == "ehf":
        return f"💡 1C o'tkazishda: oldin berilgan avans ({adv}) bo'lsa — avtomatik hisobga olinadi, qolgani {acc} qarzi"
    if row.get("direction") == "kirim":
        return f"💡 1C o'tkazishda: avval {acc} dagi xaridor qarzi yopiladi, ortgani {adv} (olingan avans)"
    return f"💡 1C o'tkazishda: avval {acc} dagi qarzimiz yopiladi, ortgani {adv} (berilgan avans)"


async def analyze_bank_row(r: BankRow, onec: OneCService, prev: dict | None = None,
                           in_file_duplicate: bool = False, extra_blockers: list[str] | tuple = (),
                           org_inn: str = "", rules=None) -> dict:
    entity = E_BANK_OUT if r.direction == "chiqim" else E_BANK_IN
    row = {**asdict(r), "date": r.date.isoformat(), "posted_date": (r.posted_date or r.date).isoformat(),
           "kind": "bank", "entity": entity, "blockers": list(extra_blockers),
           "posted": False, "posted_ref": None, "posted_number": None, "error": None, "learned": None}
    kind, blocker = special_kind(r, org_inn)
    row["special"] = kind
    kontr = await onec.find_counterparty(r.counterparty_inn) or await onec.find_counterparty_by_name(r.counterparty_name)
    if not kontr and kind == "bank_fee" and onec.mode == "live":
        kontr = await onec.fee_counterparty()  # buxgalter komissiyalarda ishlatgan kontragent
    row["hard_blockers"] = [b for b in extra_blockers if not b.startswith("⚪")]  # import fayliga ham to'siq
    row["counterparty_key"] = kontr["Ref_Key"] if kontr else None
    row["counterparty_1c_name"] = kontr.get("Description") if kontr else None
    if blocker:
        row["blockers"].append(blocker)
    elif not kontr and onec.mode == "file":
        pass  # 1C eksportida yo'q — import faylini yuklashda 1C kontragentni STIR bo'yicha o'zi yaratadi
    elif not kontr and onec.mode == "offline":
        pass  # ko'rish rejimi: kontragent 1C da qidirilmadi (umumiy "tekshirilmadi" belgisi yetarli)
    elif not kontr and kind == "bank_fee":
        row["blockers"].append("🏦 Bank xizmat haqi — firmaning o'zi 1C \"Контрагенты\" da yo'q; 1C da bankni kontragent "
                               "qilib (Прочее списание, 94.20) kiriting")
    elif not kontr:
        row["blockers"].append("Kontragent 1C da topilmadi (STIR bo'yicha) — avval katalogga qo'shing")
        row["missing_counterparty"] = True  # "➕ 1C da kontragent yaratish" tugmasi

    if in_file_duplicate:
        dup = {"status": "duplicate", "kind": "infile", "reason": "Shu tranzaksiya faylda ikki marta uchraydi",
               "match": None, "diffs": []}
    else:
        dup = await onec.check_duplicate(entity, r.number, r.date, r.amount, row["counterparty_key"],
                                         posted=r.posted_date)
    if onec.mode == "offline":
        dup = {"status": "unchecked", "kind": "none", "reason": "1C ga ulanib bo'lmadi — dublikat tekshirilmadi",
               "match": None, "diffs": []}
    row.update(status=dup["status"], status_reason=dup["reason"], match=dup["match"], diffs=dup["diffs"],
               dup_kind=dup.get("kind"))

    row["contract_ref"] = extract_contract(r.purpose)
    row["contract"] = await resolve_contract(row, onec)
    if row["contract"].get("blocker"):
        row["blockers"].append(row["contract"]["blocker"])

    if dup["status"] != "duplicate" or dup.get("kind") == "weak":
        row["learned"] = await onec.learn_history(entity, row["counterparty_key"], r.purpose, direction=r.direction)
    learned = row["learned"]
    match_code = await onec.account_code((dup["match"] or {}).get("account_key"))
    locked = rules.locked(r.counterparty_inn, r.direction, "bank", r.purpose) if rules else None
    # G'aznachilik: soliq kodi bo'yicha firmaning o'z amaliyoti — bitta G'aznachilik turli soliqlarni oladi, shuning
    # uchun u kontragent bo'yicha qat'iy qoidadan ustun (kalit so'z bo'yicha qat'iy qoida esa ustun qoladi)
    tax_hist = (await tax_code_history(onec, entity, row["counterparty_key"], tax_code_of(r.purpose))
                if r.direction == "chiqim" and is_tax_payee(r) else None)
    stored = rules.learned(r.counterparty_inn, r.direction, "bank", r.purpose) if rules else None
    if dup["status"] == "duplicate" and match_code:
        row["account"], row["account_source"] = match_code, "onec"
        row["operation"] = dup["match"].get("operation") or operation_for(match_code, r.direction)
    elif prev and prev.get("account_source") == "manual" and not (locked and locked.rule["account"] == prev["account"]):
        row["account"], row["account_source"] = prev["account"], "manual"
        row["operation"] = operation_for(row["account"], r.direction)
    elif locked and not (tax_hist and locked.rule["kind"] == "counterparty"):
        apply_rule(row, locked)
    elif tax_hist:
        # G'aznachilik: soliq kodi bo'yicha firmaning o'z amaliyoti (standart kod->schyot jadvalidan ustun)
        row["account"], row["account_source"], row["account_reason"] = tax_hist["account"], "history", tax_hist["reason"]
        row["operation"] = "ПеречислениеНалога"
        row["class_account"], row["class_reason"] = tax_hist["account"], tax_hist["reason"]
        # qoralama shu kodli oxirgi 1C hujjati namunasida: soliq turi (schyotni u belgilaydi), kontragent (1C da —
        # tuman soliq inspeksiyasi bo'lishi mumkin), holat maydonlari
        row["tax_template"] = {k: tax_hist[k] for k in ("template_ref", "tax_key", "tax_name", "counterparty_key")}
        if tax_hist["warning"]:
            row.setdefault("warnings", []).append(tax_hist["warning"])
        if tax_hist["counterparty_key"] and tax_hist["counterparty_key"] != row.get("counterparty_key"):
            name = await onec.ref_description(E_COUNTERPARTY, tax_hist["counterparty_key"])
            row["counterparty_key"], row["counterparty_1c_name"] = tax_hist["counterparty_key"], name
            row["blockers"] = [b for b in row["blockers"] if "Kontragent 1C da topilmadi" not in b]
            row.pop("missing_counterparty", None)
            row.setdefault("warnings", []).append(f"ℹ️ 1C da soliq to'lovlari «{name}» kontragentiga yoziladi — shunday qilinadi")
    elif classified := classify_purpose(r):
        # aniq belgi (soliq kodi, ish haqi, kredit …) 1C tarixidan ustun: bitta kontragent — turli schyotlar
        row["account"], row["account_source"], row["account_reason"] = classified["account"], "classified", classified["reason"]
        row["operation"] = operation_for(classified["account"], r.direction)
        row["class_account"], row["class_reason"] = classified["account"], classified["reason"]
    elif learned and learned.get("account_code"):
        row["account"], row["account_source"], row["operation"] = learned["account_code"], "history", learned["operation"]
        row["account_reason"] = (f"🧠 1C tarixi: shu kontragentning oxirgi hujjatlari ({learned['confidence']}) "
                                 f"→ {learned['account_code']}")
    elif stored:
        apply_rule(row, stored)
    else:
        row["account"] = rule_account(r)
        row["account_source"] = "rules"
        row["operation"] = operation_for(row["account"], r.direction)
    # ВидОперации — shu baza qabul qiladigan nom (O'zbekiston BP: ОплатаПокупателя, КомиссияБанка …)
    row["operation"] = await onec.valid_operation(entity, row.get("operation"), kind, (learned or {}).get("operation"))
    if prev and prev.get("posted"):
        row.update(posted=True, posted_ref=prev["posted_ref"], posted_number=prev["posted_number"], status="posted")
    if not row["contract"].get("source") and row.get("counterparty_key") and row["status"] != "duplicate":
        row["contract"] = await fallback_contract(row, onec)
    finalize_row(row)
    return row


async def fallback_contract(row: dict, onec: OneCService) -> dict:
    """To'lov maqsadida raqam yo'q: 1C tarixidagi shartnoma, bo'lmasa shu turdagi birinchi shartnoma."""
    from services_1c import E_CONTRACT
    key = (row.get("learned") or {}).get("contract_key") or (row.get("rule") or {}).get("contract_key")
    if key:
        name = await onec.ref_description(E_CONTRACT, key) or "shartnoma"
        return {"key": key, "source": "history", "name": name, "label": f"📄 {name} — 1C tarixidan"}
    contract = await onec.find_contract(row["counterparty_key"], row["direction"])
    if contract:
        return {"key": contract["Ref_Key"], "source": "default", "name": contract.get("Description", ""),
                "label": f"📄 {contract.get('Description', '')} — birinchi mos shartnoma (tekshiring)"}
    return {"key": None, "source": "none", "label": "📄 Shartnoma 1C da yo'q"}


def apply_rule(row: dict, hit) -> None:
    """Qoida bo'yicha schyot (RuleHit: qoida + sabab matni)."""
    rule = hit.rule
    row["account"], row["account_source"], row["account_reason"] = rule["account"], hit.source, hit.reason
    row["operation"] = rule["operation"] or (operation_for(rule["account"], row["direction"]) if row["kind"] == "bank"
                                             else row.get("doc_kind", "Товары"))
    row["rule"] = rule_brief(rule)
    row["rule_source"], row["rule_reason"] = hit.source, hit.reason


def finalize_row(row: dict) -> None:
    row["status_label"] = STATUS_LABELS.get(row["status"], row["status"])
    if row["status"] in ("new", "unchecked") and row.get("blockers") and not row.get("posted"):
        # "kiritishga tayyor" emas — avval to'siqni hal qilish kerak (masalan, kontragentni 1C ga qo'shish)
        row["status_label"] = ("⛔ To'siq: kontragent 1C da yo'q" if row.get("missing_counterparty")
                               else "⛔ To'siq bor (qo'lda ko'rib chiqing)")
    row["advance_account"] = advance_for(row)
    row["settlement_note"] = settlement_note(row)
    row["selectable"] = row["status"] == "new" and not row["blockers"] and not row["posted"]
    # "Baribir yozish" faqat farqli yoki raqamsiz-mos (weak) qatorlar uchun — to'liq dublikat hech qachon
    row["overridable"] = (not row["posted"] and not row["blockers"] and
                          (row["status"] == "discrepancy" or (row["status"] == "duplicate"
                                                              and row.get("dup_kind") in ("weak", "other_cp"))))
    if row["account_source"] in ("rule", "keyword", "classified") or (row["account_source"] == "history" and row.get("account_reason")):
        row["account_badge"] = row["account_reason"]
    elif row["account_source"] == "history":
        row["account_badge"] = f"🧠 Oldingi 1C tarixidan o'rganildi ({row['account']})"
    elif row["account_source"] == "manual":
        row["account_badge"] = "✍️ Buxgalter tanladi — 💾 eslab qolish mumkin"
    elif row["account_source"] == "onec":
        row["account_badge"] = "📄 1C dagi mavjud hujjatdan"
    else:
        row["account_badge"] = "🤖 AI taxmini (1C tarixi yo'q)"


async def analyze_statement(st: Statement, onec: OneCService, previous: list[dict] | None = None,
                            firm_blockers: list[str] | None = None, rules=None) -> list[dict]:
    prev = {p["id"]: p for p in previous or []}
    seen: set[tuple] = set()
    flags = []
    for r in st.rows:
        sig = (r.date, r.number, r.direction, round(r.amount, 2), r.counterparty_inn)
        flags.append(sig in seen)
        seen.add(sig)
    extra = []
    if st.is_sample and onec.mode == "live":
        extra.append("🧪 Test (namuna) vipiska — haqiqiy 1C ga yozilmaydi")
    if st.currency_code not in ("", "000"):
        extra.append(f"💱 Valyuta hisobi (kod {st.currency_code}) — hozircha faqat so'm (UZS) hisoblari yoziladi")
    if onec.mode == "offline":  # 1C ga ulanib bo'lmadi — faqat ko'rish (schyot: qoidalar/standart)
        extra.append("⚪ 1C bilan tekshirilmadi — ko'rish rejimi (1C ulangach qayta tekshiring)")
        extra.extend(firm_blockers or [])
        return list(await asyncio.gather(*(analyze_bank_row(r, onec, prev.get(r.id), f, extra, st.own_inn, rules)
                                           for r, f in zip(st.rows, flags))))
    org = await onec.organization()
    org_inn = str((org or {}).get(FIELDS["inn"], "") or "")
    if not org:
        extra.append("🏢 1C da tanlangan firmaning tashkiloti (STIR bo'yicha) topilmadi — firma sozlamalarini tekshiring")
    else:
        if st.own_inn and org_inn and st.own_inn != org_inn:
            extra.append(f"🏢 Vipiska STIR {st.own_inn} ga tegishli, tanlangan firma esa STIR {org_inn} — "
                         "boshqa firmaning vipiskasi bo'lishi mumkin")
        if st.own_account and onec.mode != "file" and not await onec.find_org_bank_account(st.own_account):
            extra.append(f"🏦 Vipiskadagi hisob raqami {st.own_account} tanlangan firmaning 1C dagi bank hisoblari "
                         "orasida yo'q — vipiska boshqa firmaniki bo'lishi mumkin")
    extra.extend(firm_blockers or [])
    return list(await asyncio.gather(*(analyze_bank_row(r, onec, prev.get(r.id), f, extra, org_inn, rules)
                                       for r, f in zip(st.rows, flags))))


def kpis(rows: list[dict]) -> dict:
    ready = [r for r in rows if r["selectable"]]
    cands = [r for r in rows if r["status"] in ("new", "posted")]
    learned = [r for r in cands if r["account_source"] in ("history", "rule", "keyword", "classified")]
    return {
        "total": len(rows), "ready": len(ready), "duplicates": sum(r["status"] == "duplicate" for r in rows),
        "discrepancies": sum(r["status"] == "discrepancy" for r in rows), "posted": sum(r["posted"] for r in rows),
        "blocked": sum(1 for r in rows if r["status"] == "new" and r["blockers"]),
        "learned": len(learned), "learned_base": len(cands),
        "learned_pct": round(100 * len(learned) / len(cands)) if cands else 0,
        "ready_sum": round(sum(r["amount"] for r in ready), 2),
    }


def cash_flow(rows: list[dict]) -> dict:
    days: dict[str, dict] = {}
    for r in rows:
        d = days.setdefault(r["date"], {"date": r["date"], "kirim": 0.0, "chiqim": 0.0})
        d[r["direction"]] += r["amount"]
    out = sorted(days.values(), key=lambda x: x["date"])
    for d in out:
        d["net"] = round(d["kirim"] - d["chiqim"], 2)
    tk = round(sum(d["kirim"] for d in out), 2)
    tc = round(sum(d["chiqim"] for d in out), 2)
    by_cp: dict[str, float] = {}
    for r in rows:
        by_cp[r["counterparty_name"]] = by_cp.get(r["counterparty_name"], 0) + (r["amount"] if r["direction"] == "kirim" else -r["amount"])
    top = sorted(by_cp.items(), key=lambda x: abs(x[1]), reverse=True)[:5]
    return {"days": out, "total_kirim": tk, "total_chiqim": tc, "net": round(tk - tc, 2),
            "top_counterparties": [{"name": n, "net": round(v, 2)} for n, v in top]}


# ---------------------------------------------------------------- qoralama (draft) payload
async def ensure_contract(row: dict, onec: OneCService, create_missing: bool) -> str | None:
    """"Yangi shartnoma" qatori: yozishdan oldin yana qidiriladi (boshqa birov yaratgan bo'lishi mumkin), bo'lmasa —
    yaratiladi. Ko'rib chiqishda (create_missing=False) hech narsa yaratilmaydi."""
    c = row.get("contract") or {}
    if c.get("source") != "new":
        return c.get("key")
    found = await onec.find_contract_by_ref(row["counterparty_key"], c["number"], c.get("date"))
    if found:
        return found["Ref_Key"]
    if not create_missing:
        return None
    created = await onec.create_contract(row["counterparty_key"], c["number"], c.get("date"), row["direction"])
    c.update(key=created["Ref_Key"], created=created.get("Description") or f"№{c['number']}")
    return created["Ref_Key"]


async def build_bank_payload(row: dict, onec: OneCService, own_account: str = "", create_missing: bool = False) -> dict:
    if not row.get("counterparty_key"):
        raise ODataError("Kontragent Ref_Key topilmadi")
    org = await onec.organization()
    if not org:
        raise ODataError("1C da tashkilot (Catalog_Организации) topilmadi")
    org_acc = await onec.find_org_bank_account(own_account or None)
    if not org_acc:
        raise ODataError("Tashkilotning bank hisobi (Catalog_БанковскиеСчета) topilmadi")
    learned = row.get("learned") or {}
    acc_key = learned.get("account_key") if learned.get("account_code") == row["account"] else None
    acc_key = acc_key or await onec.account_key(row["account"])
    if not acc_key:
        raise ODataError(f"{row['account']} schyoti 1C hisoblar rejasida topilmadi")
    rule = row.get("rule") or {}
    rule_fits = rule.get("account") == row["account"]
    # 1) to'lov maqsadidagi shartnoma (1C da topilgan) → 2) 1C tarixi / qoida → 3) shu turdagi birinchi shartnoma
    if (row.get("contract") or {}).get("source") == "new":
        contract_key = await ensure_contract(row, onec, create_missing)  # to'lov maqsadidagi YANGI shartnoma
    else:
        contract_key = ((row.get("contract") or {}).get("key") or learned.get("contract_key")
                        or (rule.get("contract_key") if rule_fits else None) or None)
    if not contract_key and (row.get("contract") or {}).get("source") != "new":
        contract = await onec.find_contract(row["counterparty_key"], row["direction"])
        contract_key = contract["Ref_Key"] if contract else None
    kontr_bank = await onec.find_bank_account(row["counterparty_key"], E_COUNTERPARTY, row.get("account_no") or None)
    cfi = learned.get("cash_flow_key") if learned.get("account_code") == row["account"] else None
    cfi = cfi or (rule.get("cash_flow_key") if rule_fits else None) or None
    operation = await onec.valid_operation(row["entity"], row["operation"], row.get("special"), learned.get("operation"))
    payload = {
        "Date": f"{row.get('posted_date') or row['date']}T12:00:00",  # bank o'tkazgan kun
        "КурсДокумента": 1, "КратностьДокумента": 1,  # so'm hisobi (valyuta hisoblari bloklanadi)
        "Организация_Key": org["Ref_Key"],
        "СчетОрганизации_Key": org_acc["Ref_Key"],
        FIELDS["operation"]: operation,
        **await onec.counterparty_payload(row["entity"], row["counterparty_key"]),
        FIELDS["contract"]: contract_key,
        "СчетКонтрагента_Key": kontr_bank["Ref_Key"] if kontr_bank else None,
        FIELDS["amount"]: row["amount"],
        FIELDS["in_number"]: row["number"],
        FIELDS["in_date"]: f"{row['date']}T00:00:00",
        FIELDS["purpose"]: row["purpose"],
        FIELDS["settlement_account"]: acc_key,
        FIELDS["cash_flow_item"]: cfi,
        FIELDS["payment_lines"]: [{
            "LineNumber": "1", FIELDS["contract"]: contract_key, "СуммаПлатежа": row["amount"],
            "СуммаВзаиморасчетов": row["amount"], "КурсВзаиморасчетов": 1, "КратностьВзаиморасчетов": 1,
            FIELDS["settlement_account"]: acc_key, FIELDS["cash_flow_item"]: cfi,
            # qarz bo'lsa — acc_key yopiladi, ortgani avansga; taqsimotni 1C o'tkazishda o'zi qiladi
            FIELDS["advance_account"]: await advance_key(row, onec),
            FIELDS["repayment"]: (learned.get("repayment") or REPAYMENT_AUTO) if row.get("advance_account") else None,
        }],
        "Комментарий": f"BUXGALTER AI: vipiskadan qoralama (№{row['number']}, {row['date']}). Tekshirib o'tkazing.",
    }
    if learned.get("account_code") == row["account"]:
        for k, v in (learned.get("extra") or {}).items():
            payload.setdefault(k, v)
    for k, v in (await onec.common_fields(row["entity"])).items():
        payload.setdefault(k, v)
    payload[FIELDS["payment_lines"]][0] = {k: v for k, v in payload[FIELDS["payment_lines"]][0].items() if v is not None}
    await mirror_template(row, payload, onec, operation, acc_key)
    return payload


async def mirror_template(row: dict, payload: dict, onec: OneCService, operation: str | None, acc_key: str | None) -> None:
    """Qoralamani 1C dagi shu turdagi oxirgi haqiqiy hujjatdek tuzish — 1C provodkani shu maydonlarga qarab
    tuzgani uchun natija bir xil chiqadi: qatorsiz ishlatiladigan operatsiyada (komissiya, soliq) — qatorsiz,
    holat maydonlari, soliq turi/davri, soliq kontragenti."""
    from services_1c import TEMPLATE_HEADER_FIELDS, _has_lines, _is_empty_key
    tax = row.get("tax_template") or {}
    tmpl = None
    if tax.get("template_ref"):
        tmpl = await onec.t.get(row["entity"], tax["template_ref"])
    if not tmpl:
        tmpl = await onec.template_doc(row["entity"], operation, row.get("counterparty_key"))
    if not tmpl:
        return
    same_cp = OneCService.counterparty_of(tmpl) == row.get("counterparty_key")
    for k in TEMPLATE_HEADER_FIELDS:
        if tmpl.get(k) not in (None, ""):
            payload[k] = tmpl[k]
    if tax:
        payload["Налог_Key"] = tax.get("tax_key") or tmpl.get("Налог_Key")
        d = date.fromisoformat(row["date"]).replace(day=1) - timedelta(days=1)  # soliq davri — o'tgan oy
        payload["НалоговыйПериод"] = f"{d.replace(day=1).isoformat()}T00:00:00"
    if same_cp and not _is_empty_key(tmpl.get("СчетКонтрагента_Key")) and (tax or not payload.get("СчетКонтрагента_Key")):
        payload["СчетКонтрагента_Key"] = tmpl["СчетКонтрагента_Key"]  # soliq organining hisobi — 1C dagidek
    for k in ("КурсДокумента", "КратностьДокумента"):  # 1C dagi qiymat va turi (masalan, "1")
        if tmpl.get(k) not in (None, "", 0):
            payload[k] = tmpl[k]
    if tax:  # soliq hujjatidagi tahliliy maydonlar (Субконто: soliq turi) — provodka analitikasi shundan
        for k, v in tmpl.items():
            if k.startswith("СубконтоДт") and "@" not in k and v not in (None, ""):
                payload[k] = v
    if _has_lines(tmpl):  # provodkani "Расшифровка платежа" qatori belgilaydi
        header = tmpl.get(FIELDS["settlement_account"])
        line_acc = (tmpl[FIELDS["payment_lines"]][0] or {}).get(FIELDS["settlement_account"])
        if header != line_acc:  # bu bazada sarlavha qatordan mustaqil (masalan, doim 42.20 yoki bo'sh) — 1C dagidek
            payload[FIELDS["settlement_account"]] = None if _is_empty_key(header) else header
        for k in ("КурсВзаиморасчетов", "КратностьВзаиморасчетов"):
            v = (tmpl[FIELDS["payment_lines"]][0] or {}).get(k)
            if v not in (None, "", 0) and payload.get(FIELDS["payment_lines"]):
                payload[FIELDS["payment_lines"]][0][k] = v
    if not _has_lines(tmpl):  # bu operatsiya 1C da "Расшифровка платежа"siz yuritiladi
        payload.pop(FIELDS["payment_lines"], None)
        tmpl_contract = tmpl.get(FIELDS["contract"])
        payload[FIELDS["contract"]] = (tmpl_contract if same_cp and not _is_empty_key(tmpl_contract)
                                       else None if _is_empty_key(tmpl_contract) else payload.get(FIELDS["contract"]))
        if not _is_empty_key(tmpl.get(FIELDS["settlement_account"])):
            payload[FIELDS["settlement_account"]] = tmpl[FIELDS["settlement_account"]]


async def advance_key(row: dict, onec: OneCService) -> str | None:
    adv = row.get("advance_account")
    if not adv:
        return None
    learned = row.get("learned") or {}
    if learned.get("advance_code") == adv and learned.get("advance_key"):
        return learned["advance_key"]
    return await onec.account_key(adv)  # 1C hisoblar rejasida topilmasa — None (1C o'z standartini qo'yadi)


async def recheck_bank(row: dict, onec: OneCService) -> dict:
    return await onec.check_duplicate(row["entity"], row["number"], date.fromisoformat(row["date"]), row["amount"],
                                      row.get("counterparty_key"),
                                      posted=date.fromisoformat(row["posted_date"]) if row.get("posted_date") else None)


def set_account(row: dict, account: str) -> None:
    from services_1c import ACCOUNT_RE
    if not ACCOUNT_RE.match(account or ""):
        raise ValueError(f"Schyot formati 60.10 (yoki 45.30.1) ko'rinishida bo'lsin: {account}")
    acc = fmt_account(account)
    learned = row.get("learned") or {}
    rule = row.get("rule") or {}
    row["account"] = acc
    if row.get("class_account") == acc:
        row["account_source"], row["account_reason"] = "classified", row["class_reason"]
        row["operation"] = operation_for(acc, row["direction"])
    elif rule.get("account") == acc and row.get("rule_source"):
        row["account_source"], row["account_reason"] = row["rule_source"], row["rule_reason"]
        row["operation"] = rule["operation"] or operation_for(acc, row["direction"])
    elif learned.get("account_code") == acc:
        row["account_source"], row["operation"] = "history", learned["operation"]
        row["account_reason"] = f"🧠 1C tarixi: shu kontragentning oxirgi hujjatlari ({learned['confidence']}) → {acc}"
    else:
        row["account_source"] = "manual"
        if row["kind"] == "bank":
            row["operation"] = operation_for(acc, row["direction"])
    finalize_row(row)


async def post_rows(rows: list[dict], ids: list[str], onec: OneCService, builders: dict,
                    force: set[str] | None = None, force_reason: str = "") -> list[dict]:
    """Tanlangan qatorlarni 1C ga QORALAMA sifatida yozadi. Har biridan oldin dublikat qayta tekshiriladi.

    builders: {"bank": (payload_builder, rechecker), "ehf": (...)}
    force: buxgalter "baribir yozish"ni alohida tasdiqlagan qatorlar. Faqat "farq bor" yoki "raqamsiz mos kelgan"
           (weak) qatorlar uchun ishlaydi; to'liq dublikat (exact) va to'siqli qatorlar HECH QACHON yozilmaydi.
    """
    force = force or set()
    results = []
    by_id = {r["id"]: r for r in rows}
    onec.invalidate()
    for rid in dict.fromkeys([*ids, *force]):
        row = by_id.get(rid)
        if not row:
            results.append({"id": rid, "ok": False, "message": "Qator topilmadi"})
            continue
        if row["posted"]:
            results.append({"id": rid, "ok": False, "message": "Allaqachon 1C ga yozilgan"})
            continue
        forced = rid in force and row.get("overridable")
        if rid in force and not forced:
            why = row["blockers"][0] if row["blockers"] else row["status_label"]
            results.append({"id": rid, "ok": False, "message": f"Majburan yozib bo'lmaydi: {why}"})
            continue
        if not row["selectable"] and not forced:
            why = row["blockers"][0] if row["blockers"] else row["status_label"]
            results.append({"id": rid, "ok": False, "message": f"Bloklandi: {why}"})
            continue
        build, recheck = builders[row["kind"]]
        try:
            fresh = await recheck(row, onec)
            still_ok = fresh["status"] == "new" or (forced and fresh.get("kind") in ("weak", "number", "other_cp"))
            if not still_ok:
                row.update(status=fresh["status"], status_reason=fresh["reason"], match=fresh["match"],
                           diffs=fresh["diffs"], dup_kind=fresh.get("kind"))
                finalize_row(row)
                results.append({"id": rid, "ok": False, "message": f"Bloklandi: {row['status_label']}"})
                continue
            payload = await build(row, onec)
            if forced and fresh["status"] != "new":
                note = f" MAJBURAN KIRITILDI ({fresh['reason']}). Buxgalter izohi: {force_reason or '—'}"
                payload["Комментарий"] = (payload.get("Комментарий", "") + note)[:1000]
            created = await onec.create_draft(row["entity"], payload)
        except ODataError as e:
            row["error"] = str(e)
            results.append({"id": rid, "ok": False, "message": str(e)})
            continue
        row.update(posted=True, posted_ref=created["ref"], posted_number=created["number"], status="posted",
                   error=None, forced=bool(forced and fresh["status"] != "new"))
        finalize_row(row)
        msg = f"Qoralama yaratildi: {created['number']}" + (" (majburan)" if row["forced"] else "")
        if (row.get("contract") or {}).get("created"):
            msg += f" · yangi shartnoma yaratildi: {row['contract']['created']}"
        results.append({"id": rid, "ok": True, "message": msg, "number": created["number"]})
    return results
