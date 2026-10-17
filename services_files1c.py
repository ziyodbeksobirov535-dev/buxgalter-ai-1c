"""1C bilan FAYL orqali ishlash (OData ishlamaganda — masalan, bulutli 1C:Fresh).

O'qish:  1C dan Excel ga chiqarilgan ro'yxat va hisobotlar ("Еще → Вывести список", hisobot → "Сохранить как"):
           • Контрагенты (Наименование, ИНН)
           • Договоры контрагентов (Наименование, Контрагент, Номер, Дата, Вид договора)
           • Банковские выписки jurnali (Дата, Номер, Поступление/Списание, Контрагент, Вид операции, …)
           • Карточка счета 51 / Журнал проводок (Дт/Кт schyotlar — qaysi to'lov qaysi schyotga olingani)
           • Оборотно-сальдовая ведомость 40.10 / 60.10 kontragentlar bo'yicha (qarz/avans qoldiqlari)
         Ulardan firmaning 1C nusxasi (FileOData) tuziladi — dublikat tekshiruvi, shartnoma va schyotni o'rganish
         xuddi OData'dagidek ishlaydi, faqat 1C ga hech narsa yozilmaydi.
Yozish:  yangi (1C da yo'q) vipiska qatorlaridan 1CClientBankExchange fayli — 1C da
         "Банк и касса → Банковские выписки → Загрузить" orqali buxgalterning o'zi yuklaydi.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from datetime import date, datetime
from pathlib import Path

import config
from accounts_uz import UZ_ACCOUNTS
from services_1c import (E_ACCOUNTS, E_BANK_ACCOUNT, E_BANK_IN, E_BANK_OUT, E_CASH_FLOW, E_CONTRACT, E_COUNTERPARTY,
                         E_NOMENCLATURE, E_ORG, E_PURCHASE, FIELDS, DemoOData, ODataError, OneCService)
from services_vipiska import clean_code, operation_for, parse_amount, parse_date

FILES_DIR = config.DATA_DIR / "onec_files"
BANK_ACCOUNT_PREFIX = "51"  # hisob-kitob schyoti (51.10) — karta/jurnaldan bank hujjatlarini ajratish uchun
KIND_LABELS = {
    "counterparties": "Kontragentlar", "contracts": "Shartnomalar", "bank_docs": "Bank hujjatlari jurnali",
    "postings": "Schyot kartochkasi / provodkalar", "osv": "Oborot-saldo vedomosti (qoldiqlar)",
}
NS = uuid.UUID("6f1c0de5-0000-4000-8000-000000000001")


class ExportFileError(ValueError):
    pass


# ---------------------------------------------------------------- nomlar
_LEGAL = (r"ооо|оoo|ooo|mchj|мчж|mas'uliyati cheklangan jamiyati?|ао|оао|зао|ап|ачж|aj|аж|xk|хк|qk|ққ|кк|чп|чф|ип|"
          r"yatt|ятт|ytt|dk|дк|duk|дук|llc|ltd|ooo|tov|тов")


def name_key(name: str) -> str:
    """'ООО "ALFA SAVDO"', 'Alfa Savdo MChJ' -> 'alfa savdo' (tashkiliy shakl, qo'shtirnoq va belgilarsiz)."""
    s = (name or "").lower().replace("ё", "е")
    s = re.sub(r"[\"'«»“”„`ʻʼ‘’]", " ", s)
    s = re.sub(r"\b(" + _LEGAL + r")\b", " ", s)
    s = re.sub(r"[^\w]+", " ", s)
    return " ".join(s.split())


def ref(*parts) -> str:
    return str(uuid.uuid5(NS, "|".join(str(p) for p in parts)))


def norm_code(v) -> str:
    """'60.10', '6010', 6010.0, '60.10 Счета к оплате' -> '60.10' ('' — schyot emas)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    m = re.match(r"^\s*(\d{2})\.?(\d{2})\b", str(v))
    return f"{m.group(1)}.{m.group(2)}" if m else ""


# ---------------------------------------------------------------- jadvalni o'qish
def _tables(filename: str, content: bytes) -> list[list[list]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except Exception as e:  # noqa: BLE001
            raise ExportFileError(f"Excel faylni o'qib bo'lmadi ({e})") from e
        return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    if name.endswith((".xls", ".mxl", ".ods")):
        raise ExportFileError("Bu formatni o'qib bo'lmaydi — 1C da \"Сохранить как\" → \"Лист Excel 2007 (*.xlsx)\" "
                              "ni tanlang")
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ExportFileError("Fayl kodirovkasi aniqlanmadi — .xlsx ga saqlang")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        class dialect(csv.excel):  # noqa: N801
            delimiter = ";"
    return [[list(r) for r in csv.reader(io.StringIO(text), dialect)]]


def _h(v) -> str:
    return " ".join(str(v or "").replace("\n", " ").lower().replace("ё", "е").split()).strip(" :")


HEADERS = [  # (maydon, regex) — tartib muhim: aniqroq birinchi
    ("in_number", r"^(вх\.?\s*(номер|№)|номер\s*вх|входящий\s*номер|номер\s*входящего|вх\.?\s*№)"),
    ("in_date", r"^(вх\.?\s*дата|дата\s*вх|дата\s*входящего)"),
    ("income", r"^(поступление|приход|сумма\s*поступления|поступило)$"),
    ("outcome", r"^(списание|расход|сумма\s*списания|списано)$"),
    ("contract_kind", r"^вид\s*договора"),
    ("operation", r"^вид\s*операции"),
    ("doc_type", r"^(вид|тип)\s*документа"),
    ("purpose", r"^назначение"),
    ("inn", r"^(инн|stir|пинфл|инн\s*/\s*пинфл|инн\s*контрагента)$"),
    ("counterparty", r"^(контрагент|владелец|контрагенты)$"),
    ("contract", r"^(договор|договор\s*контрагента)$"),
    ("account", r"^(сч[е]т\s*(учета\s*)?расч[е]тов.*|корр?\.?\s*сч[е]т|корсч[е]т|сч[е]т\s*расч[е]тов)$"),
    ("dt_account", r"^(сч[е]т\s*дт|дт\s*сч[е]т|сч[е]т\s*дебета)$"),
    ("kt_account", r"^(сч[е]т\s*кт|кт\s*сч[е]т|сч[е]т\s*кредита)$"),
    ("dt_analytics", r"^(аналитика|субконто)\s*дт"),
    ("kt_analytics", r"^(аналитика|субконто)\s*кт"),
    ("document", r"^(документ|регистратор)$"),
    ("period", r"^период$"),
    ("number", r"^(номер|№|номер\s*документа|№\s*док\.?|номер\s*договора)$"),
    ("date", r"^(дата|дата\s*документа|дата\s*договора)$"),
    ("amount", r"^(сумма|сумма\s*документа)$"),
    ("name", r"^(наименование|полное\s*наименование|рабочее\s*наименование|наименование\s*полное)$"),
    ("start", r"^сальдо\s*на\s*начало"),
    ("turnover", r"^обороты"),
    ("end", r"^сальдо\s*на\s*конец"),
    ("debit", r"^(дебет|дт)$"),
    ("credit", r"^(кредит|кт)$"),
    ("sub_account", r"^сч[е]т$"),
    ("organization", r"^организация$"),
]
HEADERS = [(f, re.compile(p)) for f, p in HEADERS]
GROUP_PARENTS = {"debit", "credit", "start", "turnover", "end"}
SUB_FIELDS = {"debit", "credit", "sub_account", "amount"}


def classify(h) -> str | None:
    t = _h(h)
    if not t:
        return None
    for f, rx in HEADERS:
        if rx.search(t):
            return f
    return None


def _find_header(table: list[list]) -> tuple[int, dict[str, int], int]:
    """Sarlavha qatori (va ikki qavatli bo'lsa — ikkinchi qator) -> {maydon: ustun}."""
    for i, row in enumerate(table[:40]):
        fields = [classify(c) for c in row]
        strong = [f for f in fields if f and f not in GROUP_PARENTS | SUB_FIELDS]
        balance_cols = sum(f in ("start", "turnover", "end") for f in fields)
        if len(strong) < 2 and not balance_cols and not (strong and any(f in GROUP_PARENTS for f in fields)):
            continue
        cols: dict[str, int] = {}
        sub = table[i + 1] if i + 1 < len(table) else []
        sub_fields = [classify(c) for c in sub]
        two_level = sum(1 for f in sub_fields if f in SUB_FIELDS) >= 2
        parent = None
        for j in range(max(len(row), len(sub) if two_level else 0)):
            f = fields[j] if j < len(fields) else None
            if f in GROUP_PARENTS:
                parent = f
            elif f is not None:
                parent = None
            s = sub_fields[j] if two_level and j < len(sub_fields) else None
            if two_level and parent and s:
                key = {("debit", "sub_account"): "dt_account", ("debit", "amount"): "dt_sum",
                       ("credit", "sub_account"): "kt_account", ("credit", "amount"): "kt_sum",
                       ("end", "debit"): "end_dt", ("end", "credit"): "end_kt",
                       ("start", "debit"): "start_dt", ("start", "credit"): "start_kt",
                       ("turnover", "debit"): "turn_dt", ("turnover", "credit"): "turn_kt"}.get((parent, s))
                if key:
                    cols.setdefault(key, j)
                continue
            if f and f not in GROUP_PARENTS:
                cols.setdefault(f, j)
        return i + (2 if two_level else 1), cols, i
    raise ExportFileError("Sarlavha qatori topilmadi (Наименование, ИНН, Дата, Сумма … ustunlari bo'lishi kerak)")


def detect_kind(cols: dict[str, int]) -> str:
    has = cols.__contains__
    if has("end_dt") or has("end_kt"):
        return "osv"
    if (has("dt_account") and has("kt_account")) and (has("dt_sum") or has("kt_sum") or has("amount")):
        return "postings"
    if has("date") and (has("income") or has("outcome") or (has("amount") and (has("counterparty") or has("purpose")
                                                                                 or has("operation")))):
        return "bank_docs"
    if has("contract_kind") or (has("counterparty") and has("number") and not has("amount")):
        return "contracts"
    if has("inn") and (has("name") or has("counterparty")):
        return "counterparties"
    raise ExportFileError("Fayl turi aniqlanmadi. Kontragentlar, shartnomalar, bank hujjatlari jurnali, "
                          "51-schyot kartochkasi yoki 40.10/60.10 oborot-saldo vedomostini yuklang")


def _cell(row: list, cols: dict, f: str):
    j = cols.get(f)
    return row[j] if j is not None and j < len(row) else None


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


DOC_RE = re.compile(r"^(.*?)\s+([\w/-]*\d[\w/-]*)\s+от\s+(\d{2}\.\d{2}\.\d{4})", re.IGNORECASE)


def parse_document_cell(v) -> dict:
    """'Списание с расчетного счета 0000-000123 от 15.10.2026 12:00:00' -> {type, number, date}."""
    m = DOC_RE.search(_text(v).replace("\n", " "))
    if not m:
        return {"type": _text(v), "number": "", "date": None}
    d = parse_date(m.group(3))
    return {"type": m.group(1).strip(), "number": m.group(2), "date": d.isoformat() if d else None}


def _direction_of(*texts) -> str:
    low = " ".join(str(t or "") for t in texts).lower()
    if "поступлен" in low or "оплата от покупателя" in low:
        return "kirim"
    if "списани" in low or "оплата поставщику" in low or "перечислени" in low:
        return "chiqim"
    return ""


def operation_code(text: str) -> str:
    """'Оплата от покупателя' -> 'ОплатаОтПокупателя' (1C ВидОперации qiymati)."""
    words = re.findall(r"[\wА-Яа-яЁё]+", text or "")
    return "".join(w[:1].upper() + w[1:].lower() for w in words)


CONTRACT_NO = re.compile(r"(?:№|N_?|#)\s*([\w/-]*\d[\w/-]*)|(?:договор\w*|shartnoma\w*)\s+([\w/-]*\d[\w/-]*)",
                         re.IGNORECASE)


def contract_number(name: str) -> str:
    m = CONTRACT_NO.search(name or "")
    return (m.group(1) or m.group(2)) if m else ""


def contract_kind(text: str) -> str:
    low = (text or "").lower()
    if "поставщик" in low or "yetkazib" in low:
        return "СПоставщиком"
    if "покупател" in low or "xaridor" in low:
        return "СПокупателем"
    return operation_code(text) or "Прочее"


# ---------------------------------------------------------------- turlar bo'yicha o'qish
def _rows(table, start, cols):
    for row in table[start:]:
        if not row or all(c in (None, "") for c in row):
            continue
        first = _h(next((c for c in row if c not in (None, "")), ""))
        if first.startswith(("итого", "всего", "jami")):
            continue
        yield row


def _counterparties(table, start, cols) -> list[dict]:
    out = []
    for row in _rows(table, start, cols):
        name = _text(_cell(row, cols, "name") or _cell(row, cols, "counterparty"))
        inn = clean_code(_cell(row, cols, "inn"))
        if name and (len(inn) in (9, 14) or not inn):
            out.append({"name": name, "inn": inn if len(inn) in (9, 14) else ""})
    return out


def _contracts(table, start, cols) -> list[dict]:
    out = []
    for row in _rows(table, start, cols):
        name = _text(_cell(row, cols, "name"))
        owner = _text(_cell(row, cols, "counterparty"))
        number = _text(_cell(row, cols, "number")) or contract_number(name)
        d = parse_date(_cell(row, cols, "date"))
        if not owner or not (name or number):
            continue
        out.append({"name": name or f"№{number}", "owner": owner, "owner_inn": clean_code(_cell(row, cols, "inn")),
                    "number": number, "date": d.isoformat() if d else None,
                    "kind": contract_kind(_text(_cell(row, cols, "contract_kind")))})
    return out


def _bank_docs(table, start, cols) -> tuple[list[dict], list[str]]:
    out, warnings = [], []
    for n, row in enumerate(_rows(table, start, cols), start=start + 1):
        d = parse_date(_cell(row, cols, "date"))
        if not d:
            continue
        inc, outc = parse_amount(_cell(row, cols, "income")) or 0, parse_amount(_cell(row, cols, "outcome")) or 0
        doc_type, op = _text(_cell(row, cols, "doc_type")), _text(_cell(row, cols, "operation"))
        if inc or outc:
            direction, amount = ("kirim", inc) if inc else ("chiqim", outc)
        else:
            amount = parse_amount(_cell(row, cols, "amount")) or 0
            direction = _direction_of(doc_type, op)
        if not amount or not direction:
            warnings.append(f"{n}-qator: summa yoki yo'nalish (kirim/chiqim) aniqlanmadi — o'tkazib yuborildi")
            continue
        in_date = parse_date(_cell(row, cols, "in_date"))
        out.append({"direction": direction, "date": d.isoformat(), "number": _text(_cell(row, cols, "number")),
                    "in_number": _text(_cell(row, cols, "in_number")),
                    "in_date": in_date.isoformat() if in_date else None, "amount": round(abs(amount), 2),
                    "counterparty": _text(_cell(row, cols, "counterparty")).split("\n")[0],
                    "inn": clean_code(_cell(row, cols, "inn")), "contract": _text(_cell(row, cols, "contract")),
                    "operation": operation_code(op) if op else "", "purpose": _text(_cell(row, cols, "purpose")),
                    "account": norm_code(_cell(row, cols, "account")), "accounts": {}})
    return out, warnings


def _postings(table, start, cols) -> tuple[list[dict], list[str]]:
    """Карточка счета 51 / Журнал проводок: bank schyoti (51.xx) bilan bog'liq provodkalar -> bank hujjatlari.
    Bitta hujjat bir nechta provodka bo'lishi mumkin (qarz 40.10 + avans 63.10) — hujjat bo'yicha jamlanadi."""
    docs: dict[tuple, dict] = {}
    last_date = None
    for row in _rows(table, start, cols):
        dt, kt = norm_code(_cell(row, cols, "dt_account")), norm_code(_cell(row, cols, "kt_account"))
        bank_dt, bank_kt = dt.startswith(BANK_ACCOUNT_PREFIX), kt.startswith(BANK_ACCOUNT_PREFIX)
        if bank_dt == bank_kt:  # bank bilan bog'liq emas yoki o'z hisoblari orasida
            continue
        amount = (parse_amount(_cell(row, cols, "dt_sum")) or parse_amount(_cell(row, cols, "kt_sum"))
                  or parse_amount(_cell(row, cols, "amount")) or 0)
        if not amount:
            continue
        direction, corr = ("kirim", kt) if bank_dt else ("chiqim", dt)
        analytics = _text(_cell(row, cols, "kt_analytics" if bank_dt else "dt_analytics"))
        lines = [x.strip() for x in analytics.split("\n") if x.strip()]
        doc = parse_document_cell(_cell(row, cols, "document"))
        d = parse_date(_cell(row, cols, "period") or _cell(row, cols, "date")) or \
            (date.fromisoformat(doc["date"]) if doc["date"] else None) or last_date
        if not d:
            continue
        last_date = d
        key = (direction, doc["number"], d.isoformat()) if doc["number"] else \
              (direction, "", d.isoformat(), round(amount, 2), name_key(lines[0] if lines else ""))
        rec = docs.setdefault(key, {
            "direction": direction, "date": d.isoformat(), "number": doc["number"], "in_number": "", "in_date": None,
            "amount": 0.0, "counterparty": lines[0] if lines else "", "inn": "",
            "contract": lines[1] if len(lines) > 1 else "", "operation": "", "purpose": "", "account": "",
            "accounts": {}})
        rec["amount"] = round(rec["amount"] + abs(amount), 2)
        rec["accounts"][corr] = round(rec["accounts"].get(corr, 0) + abs(amount), 2)
    for rec in docs.values():
        rec["account"] = main_account(rec["accounts"])
    return list(docs.values()), []


def main_account(parts: dict[str, float]) -> str:
    """Hujjat bir nechta schyotga bo'lingan bo'lsa (qarz + avans) — hisob-kitob schyoti asosiy hisoblanadi."""
    if not parts:
        return ""
    for acc in ("40.10", "60.10"):
        if acc in parts:
            return acc
    return max(parts.items(), key=lambda kv: kv[1])[0]


def _osv(table, start, cols, header_row) -> tuple[dict, list[str]]:
    """OSV (kontragentlar bo'yicha): {schyot: {nom_kaliti: {name, dt, kt}}} — davr oxiridagi qoldiq."""
    as_of = None
    for row in table[:header_row]:
        for c in row:
            for m in re.findall(r"\d{2}\.\d{2}\.\d{4}", _text(c)):
                as_of = parse_date(m) or as_of
    name_col = min((j for f, j in cols.items() if f in ("sub_account", "counterparty", "name", "account")),
                   default=0)
    out: dict[str, dict] = {}
    current = ""
    for row in _rows(table, start, cols):
        label = _text(row[name_col] if name_col < len(row) else "")
        code = norm_code(label) if re.match(r"^\s*\d{2}\.?\d{2}\b", label) else ""
        if code:
            current = code
            continue
        if not current or not label or re.search(r"договор|shartnoma|^№|^основной|документ|счет-фактура", label,
                                                 re.IGNORECASE):
            continue
        dt = parse_amount(_cell(row, cols, "end_dt")) or 0
        kt = parse_amount(_cell(row, cols, "end_kt")) or 0
        k = name_key(label)
        if k:
            cur = out.setdefault(current, {}).setdefault(k, {"name": label, "dt": 0.0, "kt": 0.0})
            cur["dt"], cur["kt"] = round(cur["dt"] + dt, 2), round(cur["kt"] + kt, 2)
    return {"as_of": as_of.isoformat() if as_of else None, "accounts": out}, []


def parse_export(filename: str, content: bytes) -> dict:
    """Bitta eksport fayli -> {"kind", "records" | "balances", "warnings"}."""
    last: Exception | None = None
    for table in _tables(filename, content):
        try:
            start, cols, header_row = _find_header(table)
            kind = detect_kind(cols)
        except ExportFileError as e:
            last = e
            continue
        warnings: list[str] = []
        if kind == "counterparties":
            records = _counterparties(table, start, cols)
        elif kind == "contracts":
            records = _contracts(table, start, cols)
        elif kind == "bank_docs":
            records, warnings = _bank_docs(table, start, cols)
        elif kind == "postings":
            records, warnings = _postings(table, start, cols)
        else:
            balances, warnings = _osv(table, start, cols, header_row)
            n = sum(len(v) for v in balances["accounts"].values())
            if not n:
                raise ExportFileError("Oborot-saldo vedomostida kontragentlar bo'yicha qoldiq topilmadi — hisobotni "
                                      "\"Контрагенты\" bo'yicha guruhlab saqlang")
            return {"kind": kind, "balances": balances, "count": n, "warnings": warnings}
        if not records:
            raise ExportFileError(f"\"{KIND_LABELS[kind]}\" fayli bo'sh — birorta yozuv o'qilmadi")
        return {"kind": kind, "records": records, "count": len(records), "warnings": warnings}
    raise last or ExportFileError("Faylda varaq topilmadi")


# ---------------------------------------------------------------- firma bo'yicha saqlash
def _path(firm_id: int) -> Path:
    return FILES_DIR / f"firm_{int(firm_id)}.json"


def load_store(firm_id: int) -> dict | None:
    p = _path(firm_id)
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text("utf-8"))
    except (OSError, ValueError):
        return None


def save_store(firm_id: int, store: dict) -> None:
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    store["stamp"] = hashlib.sha256(json.dumps(
        {k: store[k] for k in ("counterparties", "contracts", "docs", "balances")}, sort_keys=True,
        ensure_ascii=False).encode()).hexdigest()[:12]
    tmp = _path(firm_id).with_suffix(".tmp")
    tmp.write_text(json.dumps(store, ensure_ascii=False), "utf-8")
    tmp.replace(_path(firm_id))


def clear_store(firm_id: int) -> None:
    _path(firm_id).unlink(missing_ok=True)


def empty_store() -> dict:
    return {"counterparties": [], "contracts": [], "docs": [], "balances": {}, "files": [], "stamp": ""}


def _doc_key(d: dict) -> tuple:
    if d.get("number"):
        return (d["direction"], d["number"], d["date"])
    return (d["direction"], "", d["date"], round(d["amount"], 2), name_key(d.get("counterparty", "")))


def merge(store: dict, parsed: dict, filename: str, by: str = "") -> dict:
    """Yangi eksportni mavjud nusxaga qo'shish (takroriy yuklash dublikat yaratmaydi)."""
    kind = parsed["kind"]
    added = 0
    if kind == "counterparties":
        idx = {(c["inn"] or name_key(c["name"])): c for c in store["counterparties"]}
        for c in parsed["records"]:
            k = c["inn"] or name_key(c["name"])
            if k not in idx:
                added += 1
            idx[k] = {**idx.get(k, {}), **{f: v for f, v in c.items() if v}}
        store["counterparties"] = list(idx.values())
    elif kind == "contracts":
        idx = {(name_key(c["owner"]), c["number"] or name_key(c["name"]), c["date"]): c for c in store["contracts"]}
        for c in parsed["records"]:
            k = (name_key(c["owner"]), c["number"] or name_key(c["name"]), c["date"])
            added += k not in idx
            idx[k] = {**idx.get(k, {}), **{f: v for f, v in c.items() if v}}
        store["contracts"] = list(idx.values())
    elif kind in ("bank_docs", "postings"):
        idx = {_doc_key(d): d for d in store["docs"]}
        for d in parsed["records"]:
            k = _doc_key(d)
            old = idx.get(k)
            if old is None:
                added += 1
                idx[k] = d
                continue
            for f, v in d.items():  # jurnal + kartochka: bir hujjatning har xil ustunlari birlashadi
                if f == "accounts":
                    old["accounts"] = v or old.get("accounts", {})
                elif v and (not old.get(f) or kind == "postings" and f == "account"):
                    old[f] = v
        store["docs"] = sorted(idx.values(), key=lambda d: (d["date"], d["number"]))
    else:
        bal = parsed["balances"]
        accounts = store["balances"].get("accounts", {})
        accounts.update(bal["accounts"])  # shu schyotlar bo'yicha eng so'nggi vedomost
        store["balances"] = {"as_of": bal["as_of"] or store["balances"].get("as_of"), "accounts": accounts}
        added = parsed["count"]
    store["files"] = [f for f in store["files"] if not (f["name"] == filename and f["kind"] == kind)][-30:]
    store["files"].append({"name": filename, "kind": kind, "label": KIND_LABELS[kind], "count": parsed["count"],
                           "added": added, "warnings": parsed.get("warnings", [])[:5],
                           "loaded_at": datetime.now().isoformat(timespec="seconds"), "by": by})
    return {"kind": kind, "label": KIND_LABELS[kind], "count": parsed["count"], "added": added,
            "warnings": parsed.get("warnings", [])}


def summary(store: dict | None) -> dict:
    if not store:
        return {"loaded": False, "files": [], "counts": {}}
    docs = store["docs"]
    dates = [d["date"] for d in docs]
    return {"loaded": True, "files": store["files"], "stamp": store.get("stamp", ""),
            "counts": {"counterparties": len(store["counterparties"]), "contracts": len(store["contracts"]),
                       "docs": len(docs), "docs_with_account": sum(1 for d in docs if d.get("account")),
                       "balances": sum(len(v) for v in store["balances"].get("accounts", {}).values())},
            "period": {"from": min(dates), "to": max(dates)} if dates else None,
            "balances_as_of": store["balances"].get("as_of")}


# ---------------------------------------------------------------- fayldan 1C nusxasi (OData bilan bir xil interfeys)
class FileOData(DemoOData):
    """1C eksportlaridan tuzilgan faqat-o'qish nusxasi. Yozish — faqat import fayli orqali."""
    mode = "file"

    def __init__(self, db: dict):  # noqa: super().__init__ demo bazani yuklaydi — kerak emas
        self.db = db
        self._n = 0

    async def create(self, entity, payload) -> dict:
        raise ODataError("📁 1C eksport fayllari bo'yicha ishlanmoqda — 1C ga 📤 «1C import fayli» orqali kiriting")

    async def metadata_text(self) -> str:
        return ""


def build_db(store: dict, firm: dict) -> dict:
    db: dict[str, list[dict]] = {e: [] for e in (E_ORG, E_COUNTERPARTY, E_CONTRACT, E_BANK_ACCOUNT, E_NOMENCLATURE,
                                                 E_CASH_FLOW, E_ACCOUNTS, E_BANK_IN, E_BANK_OUT, E_PURCHASE)}
    org_key = ref("org", firm.get("inn") or firm.get("name"))
    db[E_ORG].append({"Ref_Key": org_key, "Description": firm.get("name", ""), FIELDS["inn"]: firm.get("inn", ""),
                      "DeletionMark": False})
    for acc in filter(None, (firm.get("bank_accounts") or "").split(",")):
        db[E_BANK_ACCOUNT].append({"Ref_Key": ref("bank", acc), "Owner": org_key, "НомерСчета": acc.strip(),
                                   "DeletionMark": False})
    codes = set(UZ_ACCOUNTS) | {d["account"] for d in store["docs"] if d.get("account")} | \
        {a for d in store["docs"] for a in d.get("accounts", {})}
    for code in sorted(c for c in codes if c):
        db[E_ACCOUNTS].append({"Ref_Key": ref("acc", code), "Code": code.replace(".", ""),
                               "Description": UZ_ACCOUNTS.get(code, ("", code))[1]})

    by_inn, by_name = {}, {}

    def counterparty(name: str, inn: str = "") -> str | None:
        k = name_key(name)
        if inn and inn in by_inn:
            return by_inn[inn]
        if k and k in by_name:
            return by_name[k]
        if not k and not inn:
            return None
        key = ref("cp", inn or k)
        db[E_COUNTERPARTY].append({"Ref_Key": key, "Description": name, FIELDS["inn"]: inn, "_name_key": k,
                                   "DeletionMark": False})
        if inn:
            by_inn[inn] = key
        if k:
            by_name.setdefault(k, key)
        return key

    for c in store["counterparties"]:
        counterparty(c["name"], c.get("inn", ""))
    contracts: dict[tuple, str] = {}
    for c in store["contracts"]:
        owner = counterparty(c["owner"], c.get("owner_inn", ""))
        key = ref("ct", owner, c.get("number"), c.get("date"), c["name"])
        db[E_CONTRACT].append({"Ref_Key": key, "Owner_Key": owner, "Description": c["name"], "Номер": c.get("number", ""),
                               "Дата": f"{c['date']}T00:00:00" if c.get("date") else None,
                               "ВидДоговора": c.get("kind", ""), "DeletionMark": False})
        contracts[(owner, name_key(c["name"]))] = key
    for d in store["docs"]:
        cp = counterparty(d.get("counterparty", ""), d.get("inn", ""))
        ct = None
        if cp and d.get("contract"):
            ct = contracts.get((cp, name_key(d["contract"])))
            if not ct:  # kartochkadagi shartnoma ro'yxatda yo'q — nomidan qo'shamiz
                ct = ref("ct", cp, d["contract"])
                db[E_CONTRACT].append({"Ref_Key": ct, "Owner_Key": cp, "Description": d["contract"],
                                       "Номер": contract_number(d["contract"]), "Дата": None,
                                       "ВидДоговора": "СПоставщиком" if d["direction"] == "chiqim" else "СПокупателем",
                                       "DeletionMark": False})
                contracts[(cp, name_key(d["contract"]))] = ct
        acc = d.get("account") or ""
        adv = next((a for a in d.get("accounts", {}) if a in ("63.10", "43.10") and a != acc), None)
        line = {FIELDS["contract"]: ct, FIELDS["settlement_account"]: ref("acc", acc) if acc else None}
        if adv:
            line[FIELDS["advance_account"]] = ref("acc", adv)
        db[E_BANK_IN if d["direction"] == "kirim" else E_BANK_OUT].append({
            "Ref_Key": ref("doc", *_doc_key(d)), "Number": d.get("number", ""), "Date": f"{d['date']}T00:00:00",
            FIELDS["in_number"]: d.get("in_number", ""),
            FIELDS["in_date"]: f"{d.get('in_date') or d['date']}T00:00:00",
            FIELDS["amount"]: d["amount"], FIELDS["counterparty"]: cp, FIELDS["contract"]: ct,
            FIELDS["operation"]: d.get("operation") or (operation_for(acc, d["direction"]) if acc else ""),
            FIELDS["settlement_account"]: ref("acc", acc) if acc else None,
            FIELDS["purpose"]: d.get("purpose", ""), FIELDS["payment_lines"]: [line],
            "Posted": True, "DeletionMark": False})
    return db


def file_service(store: dict, firm: dict) -> OneCService:
    return OneCService(FileOData(build_db(store, firm)), firm.get("inn", ""))


# ---------------------------------------------------------------- qarz / avans (OSV qoldiqlari bo'yicha)
def apply_balances(rows: list[dict], store: dict | None) -> None:
    """OSV dagi qoldiq bo'yicha: to'lovning qancha qismi qarzni yopadi, qanchasi avansga o'tadi (taxminiy ko'rinish).
    Vedomost sanasidan keyingi to'lovlar sana tartibida ketma-ket ayiriladi."""
    for r in rows:
        r.pop("balance_split", None)
    bal = (store or {}).get("balances") or {}
    accounts, as_of = bal.get("accounts") or {}, bal.get("as_of")
    if not accounts:
        return
    left: dict[tuple, float] = {}
    for r in sorted(rows, key=lambda x: (x.get("posted_date") or x["date"], x["id"])):
        acc = r.get("account")
        if r.get("kind") != "bank" or acc not in ("40.10", "60.10") or acc not in accounts \
                or r["status"] not in ("new", "unchecked", "posted"):
            continue
        if (r["direction"], acc) not in (("kirim", "40.10"), ("chiqim", "60.10")):
            continue
        if as_of and (r.get("posted_date") or r["date"]) <= as_of:
            continue  # vedomost sanasigacha bo'lgan to'lov — qoldiqda allaqachon hisobga olingan
        keys = [name_key(r.get("counterparty_1c_name") or ""), name_key(r.get("counterparty_name") or "")]
        k = next((x for x in keys if x and x in accounts[acc]), None)
        b = accounts[acc].get(k) if k else None
        debt0 = (b["dt"] - b["kt"]) if b and acc == "40.10" else (b["kt"] - b["dt"]) if b else 0.0
        remaining = left.setdefault((acc, k or keys[1]), max(debt0, 0.0))
        close = round(min(remaining, r["amount"]), 2)
        left[(acc, k or keys[1])] = round(remaining - close, 2)
        r["balance_split"] = {"account": acc, "as_of": as_of, "debt": round(remaining, 2), "close": close,
                              "advance": round(r["amount"] - close, 2), "found": bool(b)}


# ---------------------------------------------------------------- 1C ga import fayli (1CClientBankExchange)
_TRANSLIT = str.maketrans({"қ": "к", "Қ": "К", "ғ": "г", "Ғ": "Г", "ҳ": "х", "Ҳ": "Х", "ʻ": "'", "ʼ": "'", "‘": "'",
                           "’": "'", "“": '"', "”": '"', "«": '"', "»": '"', "№": "N", "—": "-", "–": "-"})


def _clean(v) -> str:
    return " ".join(str(v or "").translate(_TRANSLIT).replace("\r", " ").replace("\n", " ").split())


def exportable(row: dict) -> bool:
    return (row.get("kind") == "bank" and not row.get("posted") and row.get("status") in ("new", "unchecked")
            and not row.get("hard_blockers"))


def _d(iso: str) -> str:
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}"


def build_client_bank(rows: list[dict], own_account: str, own_inn: str, own_name: str, own_mfo: str = "",
                      created: datetime | None = None) -> bytes:
    """1C "Обмен с банком" (1CClientBankExchange 1.03) fayli, Windows-1251. Hisob-kitob schyotlarini fayl
    ko'tarmaydi — 1C ularni shartnoma / kontragent sozlamalaridan o'zi qo'yadi."""
    if not rows:
        raise ExportFileError("Eksport uchun qator yo'q")
    if not own_account:
        raise ExportFileError("Vipiskada firmaning hisob raqami (РасчСчет) aniqlanmadi")
    created = created or datetime.now()
    days = sorted(r.get("posted_date") or r["date"] for r in rows)
    lines = ["1CClientBankExchange", "ВерсияФормата=1.03", "Кодировка=Windows", "Отправитель=BUXGALTER AI",
             "Получатель=1С:Бухгалтерия", f"ДатаСоздания={created:%d.%m.%Y}", f"ВремяСоздания={created:%H:%M:%S}",
             f"ДатаНачала={_d(days[0])}", f"ДатаКонца={_d(days[-1])}", f"РасчСчет={own_account}",
             "Документ=Платежное поручение",
             "СекцияРасчСчет", f"ДатаНачала={_d(days[0])}", f"ДатаКонца={_d(days[-1])}", f"РасчСчет={own_account}",
             f"ВсегоПоступило={sum(r['amount'] for r in rows if r['direction'] == 'kirim'):.2f}",
             f"ВсегоСписано={sum(r['amount'] for r in rows if r['direction'] == 'chiqim'):.2f}",
             "КонецРасчСчет"]
    own = {"Счет": own_account, "": _clean(own_name), "ИНН": own_inn, "1": _clean(own_name),
           "РасчСчет": own_account, "БИК": own_mfo}
    for r in sorted(rows, key=lambda x: (x.get("posted_date") or x["date"], x["number"])):
        other = {"Счет": r.get("account_no") or "", "": _clean(r["counterparty_name"]), "ИНН": r.get("counterparty_inn") or "",
                 "1": _clean(r["counterparty_name"]), "РасчСчет": r.get("account_no") or "", "БИК": r.get("mfo") or ""}
        payer, payee = (own, other) if r["direction"] == "chiqim" else (other, own)
        day = _d(r.get("posted_date") or r["date"])
        lines += ["СекцияДокумент=Платежное поручение", f"Номер={_clean(r['number'])}", f"Дата={_d(r['date'])}",
                  f"Сумма={r['amount']:.2f}"]
        lines += [f"Плательщик{k}={v}" for k, v in payer.items()]
        lines += [f"Получатель{k}={v}" for k, v in payee.items()]
        lines.append(f"ДатаСписано={day}" if r["direction"] == "chiqim" else f"ДатаПоступило={day}")
        lines += ["ВидОплаты=01", f"НазначениеПлатежа={_clean(r.get('purpose'))}", "КонецДокумента"]
    lines.append("КонецФайла")
    return ("\r\n".join(lines) + "\r\n").encode("cp1251", errors="replace")


# ---------------------------------------------------------------- 📋 tekshiruv ro'yxati (Excel)
CHECKLIST_COLUMNS = [("Sana", 11), ("Bank o'tkazgan", 11), ("№", 10), ("Turi", 8), ("Kontragent", 32), ("STIR", 11),
                     ("Summa", 16), ("Holat", 26), ("Schyot", 9), ("Avans schyoti", 9), ("Schyot asosi", 40),
                     ("Shartnoma", 36), ("Qarz / avans", 48), ("To'siqlar", 40), ("Import fayli", 18),
                     ("To'lov maqsadi", 60)]


def checklist_xlsx(rows: list[dict], firm_name: str, source: str = "") -> bytes:
    """Har qator bo'yicha tavsiya: schyot, avans, shartnoma — 1C ga yuklangandan keyin qo'lda solishtirish uchun."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Tekshiruv"
    ws.append([f"BUXGALTER AI — tekshiruv ro'yxati: {firm_name}"])
    ws.append([f"Vipiska: {source} · tuzilgan {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([])
    ws.append([c for c, _w in CHECKLIST_COLUMNS])
    ws["A1"].font = Font(bold=True, size=13)
    head = PatternFill("solid", fgColor="E2E8F0")
    for i, (_c, width) in enumerate(CHECKLIST_COLUMNS, 1):
        cell = ws.cell(row=4, column=i)
        cell.font, cell.fill = Font(bold=True), head
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    fills = {"new": "FEF3C7", "duplicate": "DCFCE7", "discrepancy": "FEE2E2", "posted": "DBEAFE", "unchecked": "F1F5F9"}
    for r in sorted(rows, key=lambda x: (x["date"], x["number"])):
        ws.append([_d(r["date"]), _d(r.get("posted_date") or r["date"]), r["number"],
                   "Kirim" if r["direction"] == "kirim" else "Chiqim", r["counterparty_name"],
                   r.get("counterparty_inn") or "", r["amount"], r.get("status_label") or r["status"],
                   r.get("account") or "", r.get("advance_account") or "", r.get("account_badge") or "",
                   (r.get("contract") or {}).get("label") or "", r.get("settlement_note") or "",
                   "; ".join(r.get("blockers") or []), "chiqarilgan" if r.get("exported_at") else "",
                   r.get("purpose") or ""])
        n = ws.max_row
        ws.cell(row=n, column=7).number_format = "# ##0.00"
        fill = fills.get(r["status"])
        if fill:
            ws.cell(row=n, column=8).fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=len(CHECKLIST_COLUMNS)).column_letter}{max(ws.max_row, 5)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
