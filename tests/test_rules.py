"""📚 Hisob qoidalari: 1C tarixidan o'rganish, ustuvorlik, eslab qolish, Excel, rollar va firma izolyatsiyasi."""
import pytest
from fastapi.testclient import TestClient

import db
import main
from tests.conftest import DEMO_FIRM

PASSWORD = "parol12345"
H = {"X-Firm-Id": str(DEMO_FIRM)}


@pytest.fixture
def client():
    db.create_user("buh", PASSWORD, "accountant", firm_ids=[DEMO_FIRM])
    db.create_user("kuz", PASSWORD, "viewer", firm_ids=[DEMO_FIRM])
    c = TestClient(main.app, base_url="http://127.0.0.1:8000")
    c.post("/api/auth/login", json={"username": "buh", "password": PASSWORD})
    return c


def rules_by_match(client):
    return {(r["match"], r["direction"], r["doc_kind"]): r for r in client.get("/api/rules", headers=H).json()["rules"]}


def row(snap, number):
    return next(r for r in snap["rows"] if r["number"] == number)


def test_learn_from_history_builds_visible_table(client):
    res = client.post("/api/rules/learn", headers=H, json={"months": 12}).json()["summary"]
    assert res["documents"] == 30 and res["counterparties"] == 12 and res["keywords"] == 4 and res["added"] == 16
    rules = rules_by_match(client)
    orient = rules[("302345678", "chiqim", "bank")]
    assert orient["account"] == "60.10" and (orient["hits"], orient["total"]) == (3, 4) and orient["conflict"]
    assert rules[("komissiya", "chiqim", "bank")]["account"] == "94.20"


def test_locked_rule_beats_1c_history_and_explains_why(client):
    client.post("/api/rules", headers=H, json={"kind": "counterparty", "match": "304567890", "direction": "chiqim",
                                               "account": "69.90", "locked": True})
    snap = client.post("/api/vipiska/sample", headers=H).json()
    r = row(snap, "52")  # 1C tarixida 60.10, lekin qat'iy qoida 69.90
    assert r["account"] == "69.90" and r["account_source"] == "rule" and "Qat'iy" in r["account_badge"]
    assert "69.90" in snap["account_choices"]


def test_relearning_never_overwrites_a_locked_rule(client):
    client.post("/api/rules", headers=H, json={"kind": "counterparty", "match": "304567890", "direction": "chiqim",
                                               "account": "69.90", "locked": True})
    stats = client.post("/api/rules/learn", headers=H, json={"months": 12}).json()["summary"]
    rule = rules_by_match(client)[("304567890", "chiqim", "bank")]
    assert stats["kept_locked"] == 1 and rule["account"] == "69.90" and rule["locked"]
    assert rule["total"] == 3  # statistika yangilandi (tarixda 3 ta hujjat 60.10)


def test_keyword_rule_used_when_counterparty_unknown(client):
    client.post("/api/rules", headers=H, json={"kind": "keyword", "match": "ijara", "direction": "chiqim",
                                               "account": "94.20", "locked": False})
    csv = ("Sana;Hujjat №;Debet;Kredit;Kontragent;STIR;To'lov maqsadi\n"
           "16.10.2026;901;3 000 000,00;;Yangi Ijarachi MChJ;309999999;Oktyabr uchun ofis ijarasi\n").encode()
    snap = client.post("/api/vipiska/upload", headers=H, files={"file": ("v.csv", csv)}).json()
    r = snap["rows"][0]
    assert r["account"] == "94.20" and r["account_source"] == "keyword" and "ijara" in r["account_badge"]


def test_remember_correction_creates_locked_rule(client):
    snap = client.post("/api/vipiska/sample", headers=H).json()
    rid = row(snap, "131")["id"]  # FARG'ONA SAVDO, tarixda 40.10
    client.post(f"/api/rows/{rid}/account", headers=H, json={"account": "60.10"})
    res = client.post(f"/api/rows/{rid}/remember", headers=H, json={"mode": "vipiska"}).json()
    assert res["rule"]["locked"] and res["rule"]["source"] == "correction" and res["rule"]["account"] == "60.10"
    now = row(res, "131")  # darhol: qator endi qoida bo'yicha (✍️ emas)
    assert now["account"] == "60.10" and now["account_source"] == "rule" and "Qat'iy" in now["account_badge"]
    again = client.post("/api/vipiska/sample", headers=H).json()  # qayta yuklashda ham eslab qolingan
    assert row(again, "131")["account"] == "60.10" and row(again, "131")["account_source"] == "rule"


def test_excel_roundtrip(client):
    client.post("/api/rules/learn", headers=H, json={"months": 12})
    xlsx = client.get("/api/rules/export", headers=H)
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(xlsx.content))
    ws = wb.active
    for r in ws.iter_rows(min_row=2):
        if r[1].value == "302345678" and r[4].value == "Bank":   # ORIENT: ziddiyatni buxgalter hal qildi
            r[5].value, r[7].value = "69.90", "ha"
    ws.append(["Kalit so'z", "kommunal", "", "Chiqim", "Bank", "94.20", "", "ha"])
    ws.append(["Kontragent", "12345", "Xato STIR", "Chiqim", "Bank", "60.10", "", "ha"])
    buf = BytesIO()
    wb.save(buf)
    res = client.post("/api/rules/import", headers=H, files={"file": ("q.xlsx", buf.getvalue())}).json()
    assert res["saved"] == 17 and len(res["errors"]) == 1 and "9 xonali" in res["errors"][0]
    rules = rules_by_match(client)
    assert rules[("302345678", "chiqim", "bank")]["account"] == "69.90"
    assert rules[("302345678", "chiqim", "bank")]["locked"]
    assert rules[("kommunal", "chiqim", "bank")]["locked"]


def test_invalid_rules_rejected(client):
    bad = [{"kind": "counterparty", "match": "123", "account": "60.10"},
           {"kind": "keyword", "match": "ab", "account": "60.10"},
           {"kind": "counterparty", "match": "304567890", "account": "6010x"}]
    for body in bad:
        assert client.post("/api/rules", headers=H, json=body).status_code == 400


def test_viewer_sees_rules_but_cannot_change(client):
    client.post("/api/rules/learn", headers=H, json={"months": 12})
    viewer = TestClient(main.app, base_url="http://127.0.0.1:8000")
    viewer.post("/api/auth/login", json={"username": "kuz", "password": PASSWORD})
    data = viewer.get("/api/rules", headers=H).json()
    assert data["can_edit"] is False and data["summary"]["total"] == 16
    rid = data["rules"][0]["id"]
    assert viewer.put(f"/api/rules/{rid}", headers=H, json={"account": "69.90"}).status_code == 403
    assert viewer.delete(f"/api/rules/{rid}", headers=H).status_code == 403
    assert viewer.post("/api/rules/learn", headers=H, json={}).status_code == 403


def test_rules_are_per_firm(client):
    other = db.create_firm({"name": "Boshqa", "inn": "301111111", "onec_demo_mode": True})["id"]
    db.create_user("ega", PASSWORD, "owner")
    owner = TestClient(main.app, base_url="http://127.0.0.1:8000")
    owner.post("/api/auth/login", json={"username": "ega", "password": PASSWORD})
    owner.post("/api/rules", headers=H, json={"kind": "keyword", "match": "ijara", "account": "94.20"})
    assert owner.get("/api/rules", headers={"X-Firm-Id": str(other)}).json()["rules"] == []
    rid = owner.get("/api/rules", headers=H).json()["rules"][0]["id"]
    # boshqa firma sarlavhasi bilan bu firma qoidasini o'zgartirib bo'lmaydi
    assert owner.put(f"/api/rules/{rid}", headers={"X-Firm-Id": str(other)}, json={"account": "60.10"}).status_code == 400
    assert client.put(f"/api/rules/{rid}", headers={"X-Firm-Id": str(other)}, json={"account": "60.10"}).status_code == 403


def test_audit_records_rule_changes(client):
    client.post("/api/rules/learn", headers=H, json={"months": 12})
    client.post("/api/rules", headers=H, json={"kind": "keyword", "match": "ijara", "account": "94.20"})
    actions = [a["action"] for a in db.list_audit()]
    assert "rules_learned" in actions and "rule_saved" in actions
