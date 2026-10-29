"""Dastur ko'rigidan keyingi tuzatishlar: haqiqiy firmada demo ma'lumot yo'q, OData ishlamasa fayllarga o'tish,
Didox QQS stavkasi, bot/AI fayl rejimi, tekshiruv ro'yxati. Ma'lumotlar o'ylab topilgan."""
import asyncio
import io
import types
from datetime import date
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import bot
import core
import db
import main
import services_ai
from services_1c import DemoOData, ODataError, OneCService
from services_didox import DidoxError, _demo_invoices, _vat_rate, analyze_invoice, build_purchase_payload
from tests.test_files1c import ALFA, CARD_51, COUNTERPARTIES, JOURNAL, bank, load, statement

PASSWORD = "parol12345"
TG = 5151


def run(coro):
    return asyncio.run(coro)


@pytest.fixture
def files_firm():
    """1C manzili yo'q, faqat eksport fayllari yuklangan haqiqiy firma."""
    fid = db.create_firm({"name": "TEST FIRMA", "inn": "300000777", "bank_accounts": "20208000100000000777"})["id"]
    load(fid, ("c.xlsx", COUNTERPARTIES), ("j.xlsx", JOURNAL), ("k.xlsx", CARD_51))
    return fid


@pytest.fixture
def broken_live_firm(monkeypatch):
    """1C manzili bor, lekin OData hech narsa bermaydi (OData obyektlari nashr qilinmagan holat) + eksport fayllari yuklangan."""
    fid = db.create_firm({"name": "TEST FIRMA", "inn": "300000777", "onec_demo_mode": False,
                          "onec_base_url": "https://1c.example.uz/odata", "onec_user": "u", "onec_password": "p",
                          "didox_demo_mode": False, "didox_user_key": "k"})["id"]
    load(fid, ("c.xlsx", COUNTERPARTIES), ("k.xlsx", CARD_51))

    async def broken(*a, **kw):
        raise ODataError("OData: obyektlar nashr qilinmagan")
    t = core.get_onec(fid).t
    for name in ("query", "get", "ping", "metadata_text"):
        monkeypatch.setattr(t, name, broken)
    return fid


def client_for(firm_id, role="accountant", username="buh"):
    user = db.create_user(username, PASSWORD, role, "B", firm_ids=[firm_id])
    c = TestClient(main.app, base_url="http://127.0.0.1:8000")
    c.post("/api/auth/login", json={"username": username, "password": PASSWORD})
    c.user = user
    return c, {"X-Firm-Id": str(firm_id)}


# ---------------------------------------------------------------- haqiqiy firmada demo ma'lumot ko'rsatilmaydi
def test_demo_didox_invoices_are_not_shown_for_a_real_firm(files_firm):
    session = core.get_session("r1", files_firm)
    with pytest.raises(DidoxError, match="Demo hisob-fakturalar faqat Demo firmada"):
        run(core.load_didox(session))
    # Didox tanlangan, lekin hamkor kaliti yo'q — tarmoqqa chiqilmaydi, demo ham ko'rsatilmaydi
    db.update_firm(files_firm, {"didox_demo_mode": False, "didox_password": "synthetic-password"})
    with pytest.raises(DidoxError, match="hamkor kaliti"):
        run(core.load_didox(session))


def test_sample_statement_only_in_demo_firm(files_firm):
    c, h = client_for(files_firm)
    r = c.post("/api/vipiska/sample", headers=h)
    assert r.status_code == 400 and "Demo firmada" in r.json()["detail"]


def test_ai_does_not_load_sample_into_real_firm(files_firm):
    ctx = services_ai.ToolContext(core.get_session("r2", files_firm), "vipiska")
    assert "xato" in run(ctx.t_analyze_and_deduplicate_vipiska())
    assert "xato" in run(ctx.t_cash_flow_summary())


# ---------------------------------------------------------------- OData ishlamasa — eksport fayllari
def test_status_learning_and_didox_fall_back_to_exports(broken_live_firm, monkeypatch):
    st = run(core.onec_status(broken_live_firm))
    assert st["ok"] and st["mode"] == "file" and "OData javob bermadi" in st["warning"]

    c, h = client_for(broken_live_firm)
    res = c.post("/api/rules/learn", headers=h, json={"months": 24}).json()
    assert res["summary"]["source"] == "file" and res["summary"]["counterparties"] >= 1

    async def invoices(_settings):
        return _demo_invoices(), "live"
    monkeypatch.setattr(core, "fetch_invoices", invoices)
    session = core.get_session("r3", broken_live_firm)
    snap = run(core.load_didox(session))
    assert snap["files_mode"] and {r["status"] for r in snap["rows"]} == {"unchecked"}
    assert not any("Nomenklatura" in b for r in snap["rows"] for b in r["blockers"])  # eksportda nomenklatura yo'q
    with pytest.raises(ODataError):
        run(core.post(session, "didox", [r["id"] for r in snap["rows"]]))


def test_ai_post_tool_points_to_import_file_in_file_mode(files_firm):
    session = core.get_session("r4", files_firm)
    session.statement = statement(bank(date(2026, 10, 2), "78", "chiqim", 3_000_000, ALFA))
    run(core._analyze_vipiska(session))
    ctx = services_ai.ToolContext(session, "vipiska")
    res = run(ctx.t_post_confirmed_to_1c())
    assert res["holat"] == "import_fayli_orqali" and res["import_uchun_tayyor"] == 1 and not ctx.cards
    assert "import fayli" in services_ai.mock_reply([("post_confirmed_to_1c", {})], [res], "vipiska")


# ---------------------------------------------------------------- Didox QQS stavkasi
@pytest.mark.parametrize("item, rate", [({"vat_rate": 12}, 12.0), ({"vat_rate": "0"}, 0.0), ({"without_vat": True}, 0.0),
                                        ({}, None)])
def test_didox_vat_rate_mapping(item, rate):
    assert _vat_rate(item) == rate


def test_vat_free_invoice_has_no_vat_line():
    onec = OneCService(DemoOData())
    inv = next(i for i in _demo_invoices() if i["number"] == "AP-1007")
    for it in inv["items"]:
        it["vat_rate"] = 0.0  # aylanma solig'i to'lovchisidan — QQSsiz
    row = run(analyze_invoice(inv, onec))
    assert row["vat_account"] == "" and "44.10" not in row["posting_note"]
    line = run(build_purchase_payload(row, onec))["Товары"][0]
    assert line["СтавкаНДС"] == "БезНДС" and line["СуммаНДС"] == 0 and "СчетУчетаНДС_Key" not in line
    assert line["Сумма"] == inv["items"][0]["total"]


def test_vat_invoice_keeps_12_percent():
    onec = OneCService(DemoOData())
    inv = next(i for i in _demo_invoices() if i["number"] == "AP-1007")
    row = run(analyze_invoice(inv, onec))
    line = run(build_purchase_payload(row, onec))["Товары"][0]
    assert row["vat_account"] == "44.10" and line["СтавкаНДС"] == "НДС12"
    assert line["СуммаНДС"] == round(line["Сумма"] - line["Сумма"] / 1.12, 2)


# ---------------------------------------------------------------- 📋 tekshiruv ro'yxati
def test_checklist_excel_lists_accounts_contracts_and_status(files_firm):
    c, h = client_for(files_firm, role="viewer", username="kuz")  # ko'rish — kuzatuvchiga ham ruxsat
    assert c.get("/api/vipiska/checklist", headers=h).status_code == 400  # vipiska yo'q
    session = core.get_session(core.session_id(c.user["id"], files_firm), files_firm)
    session.statement = statement(bank(date(2026, 10, 2), "78", "chiqim", 3_000_000, ALFA,
                                       "Договор №45 от 01.02.2026 оплата"))
    run(core._analyze_vipiska(session))
    r = c.get("/api/vipiska/checklist", headers=h)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[4]]
    row = dict(zip(header, [c.value for c in ws[5]]))
    assert row["№"] == "78" and row["Schyot"] == "60.10" and row["Avans schyoti"] == "43.10"
    assert "Shartnoma №45" in row["Shartnoma"] and row["Summa"] == 3_000_000


# ---------------------------------------------------------------- Telegram: import fayli
def test_bot_sends_import_file_in_file_mode(files_firm):
    user = db.create_user("buh", PASSWORD, "accountant", "B", firm_ids=[files_firm])
    db.consume_link_code(db.create_link_code(user["id"]), TG)
    db.set_telegram_firm(user["id"], files_firm)
    _, firm, session = bot.context_of(TG)
    session.statement = statement(bank(date(2026, 10, 2), "78", "chiqim", 3_000_000, ALFA))
    run(core._analyze_vipiska(session))
    snap = core.snapshot(session)
    assert "eksport fayllari" in bot.summary_text(snap, firm)
    data = [b.callback_data for row in bot.result_keyboard(snap, user, session).inline_keyboard for b in row
            if b.callback_data]
    assert data == [f"export:{files_firm}"]  # to'g'ridan-to'g'ri "1C ga kiritish" tugmasi yo'q
    msg = types.SimpleNamespace(answer_document=AsyncMock(), edit_reply_markup=AsyncMock())
    cb = types.SimpleNamespace(data=data[0], from_user=types.SimpleNamespace(id=TG), message=msg, answer=AsyncMock())
    run(bot.cb_export(cb))
    sent = msg.answer_document.call_args.args[0]
    assert sent.filename.startswith(f"kl_to_1c_{files_firm}_") and b"1CClientBankExchange" in sent.data
    run(bot.cb_export(cb))  # ikkinchi marta — yangi qator yo'q
    assert "yangi qator yo'q" in cb.answer.call_args.args[0]
