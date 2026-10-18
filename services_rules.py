"""📚 Hisob qoidalari: buxgalter bazani qanday yuritganini 1C tarixidan o'rganish va qo'llash.

Ustuvorlik (vipiska qatori uchun schyot tanlash):
  1. 🔒 Qat'iy qoida — kontragent bo'yicha, keyin kalit so'z bo'yicha (buxgalter tasdiqlagan)
  2. 🧠 Shu kontragentning 1C dagi oxirgi hujjatlari (jonli so'rov)
  3. 🧠 1C tarixidan o'rganilgan kontragent qoidasi (12 oylik statistika)
  4. 🔤 Kalit so'z qoidasi (to'lov maqsadi bo'yicha)
  5. 🤖 Standart taxmin
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from services_1c import (DT, E_ACCOUNTS, E_BANK_IN, E_BANK_OUT, E_COUNTERPARTY, E_PURCHASE, FIELDS, PURPOSE_KEYWORDS,
                         Cond, OneCService, _has_lines, fmt_account)

# To'lov maqsadida qidiriladigan kalit so'zlar — tarixda shu so'z bilan kelgan to'lovlar qaysi schyotga olingan
KEYWORD_CANDIDATES = sorted({kw for kw, _acc, _dir in PURPOSE_KEYWORDS} | {
    "ijara", "аренд", "kommunal", "коммунал", "elektr", "электроэнерг", "internet", "aloqa", "связь",
    "sug'urta", "страхов", "foiz", "процент", "qarz", "займ", "avans", "аванс", "bojxona", "тамож",
    "yoqilg'i", "топлив", "benzin", "reklama", "реклам", "audit", "konsalting", "консульт", "transport",
    "транспорт", "jarima", "штраф", "pensiya", "пенсион", "dividend", "дивиденд",
})
MIN_KEYWORD_DOCS = 2
MIN_KEYWORD_SHARE = 0.7
PAGE = 200
MAX_DOCS_PER_TYPE = 5000

DIRECTION_LABELS = {"kirim": "Kirim", "chiqim": "Chiqim", "": "Har ikkisi"}
DOC_KIND_LABELS = {"bank": "Bank", "ehf": "EHF (xarid)"}


# ---------------------------------------------------------------- qoidalarni qo'llash
@dataclass
class RuleHit:
    rule: dict
    reason: str
    source: str  # rule | history | keyword


class RuleSet:
    """Bitta firmaning qoidalari — tez qidirish uchun indekslangan."""

    def __init__(self, rules: list[dict]):
        self.by_cp: dict[tuple, dict] = {}
        self.keywords: list[dict] = []
        for r in rules:
            if r["kind"] == "counterparty":
                self.by_cp[(r["match"], r["direction"], r["doc_kind"])] = r
            else:
                self.keywords.append(r)
        self.keywords.sort(key=lambda r: -len(r["match"]))  # uzunroq (aniqroq) so'z birinchi

    def _cp(self, inn: str, direction: str, doc_kind: str) -> dict | None:
        if not inn:
            return None
        return self.by_cp.get((inn, direction, doc_kind)) or self.by_cp.get((inn, "", doc_kind))

    def _kw(self, purpose: str, direction: str, doc_kind: str, locked: bool) -> dict | None:
        low = (purpose or "").lower()
        for r in self.keywords:
            if r["locked"] == locked and r["doc_kind"] == doc_kind and r["direction"] in ("", direction) \
                    and r["match"] in low:
                return r
        return None

    def locked(self, inn: str, direction: str, doc_kind: str, purpose: str) -> RuleHit | None:
        cp = self._cp(inn, direction, doc_kind)
        if cp and cp["locked"]:
            return RuleHit(cp, f"🔒 Qat'iy qoida: shu kontragent → {cp['account']}", "rule")
        kw = self._kw(purpose, direction, doc_kind, locked=True)
        if kw:
            return RuleHit(kw, f"🔒 Qat'iy qoida: «{kw['match']}» so'zi → {kw['account']}", "rule")
        return None

    def learned(self, inn: str, direction: str, doc_kind: str, purpose: str) -> RuleHit | None:
        cp = self._cp(inn, direction, doc_kind)
        if cp and not cp["locked"]:
            stat = f" ({cp['hits']}/{cp['total']} hujjat)" if cp["total"] else ""
            return RuleHit(cp, f"🧠 1C tarixidan o'rganilgan qoida{stat} → {cp['account']}", "history")
        kw = self._kw(purpose, direction, doc_kind, locked=False)
        if kw:
            return RuleHit(kw, f"🔤 Kalit so'z «{kw['match']}» → {kw['account']}", "keyword")
        return None


# ---------------------------------------------------------------- 1C tarixidan o'rganish
async def _fetch_all(onec: OneCService, entity: str, conds: list[Cond], orderby: str | None = None,
                     select: list[str] | None = None, limit: int = MAX_DOCS_PER_TYPE) -> list[dict]:
    out: list[dict] = []
    # Tartib barqaror bo'lishi shart: $orderby siz bulutli 1C $skip sahifalarini takrorlab beradi (yozuvlar yo'qoladi)
    orderby = f"{orderby},Ref_Key" if orderby else "Ref_Key"
    while len(out) < limit:
        page = await onec.t.query(entity, conds, top=PAGE, skip=len(out) or None, orderby=orderby, select=select)
        out.extend(page)
        if len(page) < PAGE:
            break
    return out[:limit]


def _mode(values) -> str:
    c = Counter(v for v in values if v)
    return c.most_common(1)[0][0] if c else ""


async def learn_from_1c(onec: OneCService, months: int = 12) -> tuple[list[dict], dict]:
    """1C dagi o'tkazilgan bank va xarid hujjatlaridan qoidalar jadvalini tuzadi (1C ga hech narsa yozmaydi)."""
    cutoff = datetime.combine(date.today() - timedelta(days=int(months * 30.5)), datetime.min.time())
    counterparties = {r["Ref_Key"]: (str(r.get(FIELDS["inn"]) or ""), r.get("Description", ""))
                      for r in await _fetch_all(onec, E_COUNTERPARTY, [Cond("DeletionMark", "eq", False)],
                                                select=["Ref_Key", FIELDS["inn"], "Description"], limit=50000)}
    accounts = {r["Ref_Key"]: fmt_account(r.get("Code")) for r in
                await _fetch_all(onec, E_ACCOUNTS, [], select=["Ref_Key", "Code"], limit=5000)}
    base = [Cond("Posted", "eq", True), Cond("DeletionMark", "eq", False), Cond("Date", "ge", DT(cutoff))]

    rules: list[dict] = []
    summary = {"months": months, "documents": 0, "counterparties": 0, "keywords": 0, "conflicts": 0,
               "skipped_no_inn": 0}
    for entity, direction, doc_kind in ((E_BANK_OUT, "chiqim", "bank"), (E_BANK_IN, "kirim", "bank"),
                                        (E_PURCHASE, "", "ehf")):
        docs = await _fetch_all(onec, entity, base, orderby="Date desc")
        summary["documents"] += len(docs)
        posted = await onec.bank_postings(entity, cutoff)  # qatorsiz hujjatlar uchun — 1C provodkasidagi schyot
        groups: dict[str, list[tuple[dict, dict]]] = defaultdict(list)
        kw_votes: dict[str, Counter] = defaultdict(Counter)
        kw_ops: dict[tuple, Counter] = defaultdict(Counter)
        for d in docs:
            p = OneCService._extract_pattern(d)
            if not _has_lines(d) and posted.get(d.get("Ref_Key")):
                p["account_key"] = posted[d["Ref_Key"]]
            code = accounts.get(p["account_key"]) if p["account_key"] else None
            if not code:
                continue
            groups[OneCService.counterparty_of(d) or ""].append((d, {**p, "code": code}))
            if doc_kind == "bank":
                low = str(d.get(FIELDS["purpose"]) or "").lower()
                for kw in KEYWORD_CANDIDATES:
                    if kw in low:
                        kw_votes[kw][code] += 1
                        kw_ops[(kw, code)][p["operation"]] += 1

        for kontr_key, items in groups.items():
            inn, name = counterparties.get(kontr_key, ("", ""))
            if len(inn) != 9:
                summary["skipped_no_inn"] += 1
                continue
            votes = Counter((p["code"], p["operation"]) for _d, p in items)
            (code, op), hits = votes.most_common(1)[0]
            best = [p for _d, p in items if (p["code"], p["operation"]) == (code, op)]
            variants = [{"account": c, "operation": o, "count": n} for (c, o), n in votes.most_common()]
            rules.append({
                "kind": "counterparty", "match": inn, "name": name, "direction": direction, "doc_kind": doc_kind,
                "account": code, "advance_account": accounts.get(_mode(p["advance_key"] for p in best), "") or "",
                "operation": op or "", "contract_key": _mode(p["contract_key"] for p in best),
                "cash_flow_key": _mode(p["cash_flow_key"] for p in best), "hits": hits, "total": len(items),
                "variants": variants, "last_date": max(str(d.get("Date", ""))[:10] for d, _p in items),
            })
            summary["counterparties"] += 1
            summary["conflicts"] += len(variants) > 1

        for kw, votes in kw_votes.items():
            total = sum(votes.values())
            code, hits = votes.most_common(1)[0]
            if total >= MIN_KEYWORD_DOCS and hits / total >= MIN_KEYWORD_SHARE:
                rules.append({
                    "kind": "keyword", "match": kw, "name": "", "direction": direction, "doc_kind": doc_kind,
                    "account": code, "operation": kw_ops[(kw, code)].most_common(1)[0][0] or "", "contract_key": "",
                    "cash_flow_key": "", "hits": hits, "total": total,
                    "variants": [{"account": c, "operation": "", "count": n} for c, n in votes.most_common()],
                    "last_date": "",
                })
                summary["keywords"] += 1
    return rules, summary


# ---------------------------------------------------------------- Excel
EXCEL_HEADERS = ["Turi", "STIR yoki kalit so'z", "Nomi", "Yo'nalish", "Hujjat", "Schyot", "Avans schyoti", "ВидОперации",
                 "Qat'iy (ha/yo'q)", "Ishonch", "Variantlar", "Manba", "Oxirgi sana"]
_SOURCE_LABELS = {"history": "1C tarixi", "manual": "Buxgalter", "correction": "Tuzatishdan", "import": "Excel"}


def export_xlsx(rules: list[dict], firm_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Qoidalar"
    ws.append(EXCEL_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in rules:
        ws.append([
            "Kontragent" if r["kind"] == "counterparty" else "Kalit so'z", r["match"], r["name"],
            DIRECTION_LABELS[r["direction"]], DOC_KIND_LABELS[r["doc_kind"]], r["account"], r["advance_account"], r["operation"],
            "ha" if r["locked"] else "yo'q", f"{r['hits']}/{r['total']}" if r["total"] else "",
            "; ".join(f"{v['account']}×{v['count']}" for v in r["variants"]), _SOURCE_LABELS.get(r["source"], r["source"]),
            r["last_date"]])
        ws.cell(row=ws.max_row, column=2).number_format = "@"
    for col, width in zip("ABCDEFGHIJKLM", (12, 18, 40, 11, 12, 9, 13, 34, 15, 9, 22, 12, 12)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Yo'riqnoma")
    for line in [f"Firma: {firm_name}", "",
                 "Tahrirlash mumkin: Schyot, Avans schyoti, ВидОперации, Qat'iy (ha/yo'q). Yangi qator ham qo'shish mumkin.",
                 "Schyot — qarz bo'yicha (40.10 / 60.10); Avans schyoti — qarz bo'lmasa (63.10 / 43.10). 1C to'lovni "
                 "o'tkazishda avval qarzni yopadi, ortganini avansga o'tkazadi.",
                 "Turi: Kontragent (STIR 9 xona) yoki Kalit so'z (to'lov maqsadida qidiriladi).",
                 "Yo'nalish: Kirim / Chiqim / Har ikkisi.  Hujjat: Bank yoki EHF (xarid).",
                 "Qat'iy = ha — bu qoida 1C tarixidan ustun turadi.",
                 "Ishonch, Variantlar, Manba, Oxirgi sana — faqat ma'lumot uchun (yuklashda e'tiborga olinmaydi)."]:
        info.append([line])
    info.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_import_xlsx(content: bytes) -> tuple[list[dict], list[str]]:
    """Excel'dan qoidalar. Natija: (qoidalar, xatolar ro'yxati qator raqami bilan)."""
    from openpyxl import load_workbook

    try:
        ws = load_workbook(io.BytesIO(content), data_only=True, read_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"Excel faylni o'qib bo'lmadi: {e}") from e
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Fayl bo'sh")
    header = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(header):
                if h.startswith(n):
                    return i
        return None
    idx = {"kind": col("turi"), "match": col("stir", "kalit"), "name": col("nomi"), "direction": col("yo'nalish", "yonalish"),
           "doc_kind": col("hujjat"), "account": col("schyot", "счет"), "advance_account": col("avans"),
           "operation": col("видоперации", "vid"),
           "locked": col("qat'iy", "qatiy")}
    if idx["kind"] is None or idx["match"] is None or idx["account"] is None:
        raise ValueError("Ustunlar topilmadi: kamida 'Turi', 'STIR yoki kalit so'z', 'Schyot' bo'lishi kerak "
                         "(avval qoidalarni Excel'ga yuklab olib, o'shani tahrirlang)")
    out, errors = [], []
    for n, row in enumerate(rows[1:], start=2):
        def get(k):
            i = idx[k]
            return row[i] if i is not None and i < len(row) else None
        if not any(v not in (None, "") for v in row):
            continue
        kind_raw = str(get("kind") or "").lower()
        kind = "keyword" if "kalit" in kind_raw or "keyword" in kind_raw else "counterparty"
        match = get("match")
        if isinstance(match, float) and match.is_integer():
            match = int(match)
        direction_raw = str(get("direction") or "").lower()
        direction = "kirim" if "kirim" in direction_raw else "chiqim" if "chiqim" in direction_raw else ""
        doc_kind = "ehf" if "ehf" in str(get("doc_kind") or "").lower() else "bank"
        locked = str(get("locked") or "").strip().lower() in ("ha", "yes", "да", "1", "true", "✓")
        account = get("account")
        if isinstance(account, float):
            account = f"{account:.2f}"
        advance = get("advance_account")
        if isinstance(advance, float):
            advance = f"{advance:.2f}"
        out.append({"_row": n, "kind": kind, "match": str(match or ""), "name": str(get("name") or ""),
                    "direction": direction, "doc_kind": doc_kind, "account": str(account or ""),
                    "advance_account": str(advance or ""),
                    "operation": str(get("operation") or ""), "locked": locked})
    return out, errors
